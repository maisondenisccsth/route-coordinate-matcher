"""
Route Coordinate Matcher — Streamlit Web App
==============================================
โยนไฟล์ route เข้าไป ได้พิกัดกลับมาทันที ผ่านเว็บแอปที่มีลิงก์ถาวร

วิธี deploy: ดู DEPLOY_INSTRUCTIONS.md
"""

import streamlit as st
import pandas as pd
import io
import re
import xlrd
from datetime import datetime
from openpyxl import load_workbook, Workbook
from openpyxl.utils import get_column_letter
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
import qrcode

st.set_page_config(page_title="Route Coordinate Matcher", page_icon="🚚", layout="wide")

# ============================================================
# THEME — Logistics / Transportation, refined
# ============================================================
st.markdown("""
<style>
    @import url('https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600;700;800&family=JetBrains+Mono:wght@400;500;600;700&display=swap');

    html, body, [class*="css"] { font-family: 'Inter', sans-serif; }

    .stApp {
        background:
            radial-gradient(1200px 600px at 90% -10%, rgba(255,107,53,0.08), transparent 60%),
            radial-gradient(1000px 500px at -10% 10%, rgba(47,128,237,0.10), transparent 60%),
            linear-gradient(180deg, #070C16 0%, #0B1524 45%, #0A1220 100%);
    }

    /* ---------- Header (with photo) ---------- */
    .rcm-header {
        background:
            linear-gradient(100deg, rgba(9,17,32,0.94) 0%, rgba(11,22,40,0.85) 40%, rgba(11,22,40,0.35) 75%, rgba(11,22,40,0.15) 100%),
            url('https://images.unsplash.com/photo-1601584115197-04ecc0da31d7?auto=format&fit=crop&w=1600&q=80');
        background-size: cover;
        background-position: center 60%;
        border: 1px solid #204A78;
        border-radius: 18px;
        padding: 30px 34px 24px 34px;
        margin-bottom: 22px;
        position: relative;
        overflow: hidden;
        box-shadow: 0 10px 40px rgba(0,0,0,0.4), inset 0 1px 0 rgba(255,255,255,0.04);
        min-height: 210px;
    }
    .rcm-header::before {
        content: "";
        position: absolute;
        left: 0; bottom: 0; width: 100%; height: 3px;
        background: linear-gradient(90deg, #FF6B35 0%, #FFB347 35%, transparent 75%);
    }
    .rcm-eyebrow {
        font-family: 'JetBrains Mono', monospace;
        font-size: 11.5px;
        letter-spacing: 0.18em;
        color: #FF9B5C;
        font-weight: 700;
        margin-bottom: 8px;
        text-transform: uppercase;
        display: flex; align-items: center; gap: 6px;
    }
    .rcm-header-content {
        position: relative;
        z-index: 2;
    }
    .rcm-route-svg {
        position: absolute;
        top: 0; right: 0;
        width: 60%; height: 100%;
        z-index: 1;
        pointer-events: none;
    }
    .rcm-title {
        font-size: 30px;
        font-weight: 800;
        color: #F7FAFD;
        margin: 0 0 6px 0;
        letter-spacing: -0.02em;
        text-shadow: 0 2px 12px rgba(0,0,0,0.5);
    }
    .rcm-subtitle {
        font-size: 14px;
        color: #C3D4EA;
        font-family: 'JetBrains Mono', monospace;
        margin-bottom: 20px;
    }
    .rcm-features {
        display: flex;
        gap: 28px;
        flex-wrap: wrap;
    }
    .rcm-feature {
        display: flex;
        align-items: center;
        gap: 8px;
    }
    .rcm-feature-icon {
        width: 30px; height: 30px;
        border-radius: 8px;
        background: rgba(255,255,255,0.08);
        border: 1px solid rgba(255,255,255,0.14);
        display: flex; align-items: center; justify-content: center;
        font-size: 15px;
        flex-shrink: 0;
    }
    .rcm-feature-text {
        line-height: 1.25;
    }
    .rcm-feature-title {
        font-size: 12.5px;
        font-weight: 700;
        color: #F0F5FA;
    }
    .rcm-feature-sub {
        font-size: 10.5px;
        color: #92A9C4;
        font-family: 'JetBrains Mono', monospace;
    }

    /* ---------- Status / connection card ---------- */
    .rcm-status-card {
        display: flex;
        align-items: center;
        justify-content: space-between;
        background: linear-gradient(90deg, rgba(42,110,76,0.20), rgba(12,23,40,0.55));
        border: 1px solid #2A6E4C;
        border-radius: 14px;
        padding: 16px 22px;
        margin-bottom: 14px;
        box-shadow: 0 4px 18px rgba(0,0,0,0.22);
    }
    .rcm-status-left { display: flex; align-items: center; gap: 14px; }
    .rcm-status-icon {
        width: 38px; height: 38px; border-radius: 10px;
        background: rgba(52,199,89,0.18);
        border: 1px solid rgba(52,199,89,0.4);
        display: flex; align-items: center; justify-content: center;
        font-size: 17px;
    }
    .rcm-status-title { font-weight: 700; color: #F0F5FA; font-size: 14.5px; }
    .rcm-status-sub { font-size: 11.5px; color: #8FA8C7; font-family: 'JetBrains Mono', monospace; }
    .rcm-status-right { display: flex; align-items: center; gap: 8px; }
    .rcm-status-right-icon {
        width: 34px; height: 34px; border-radius: 9px;
        background: rgba(255,255,255,0.06);
        border: 1px solid rgba(255,255,255,0.12);
        display: flex; align-items: center; justify-content: center;
        font-size: 15px;
    }
    .rcm-status-pill {
        font-family: 'JetBrains Mono', monospace;
        font-size: 11px; color: #6FE39A; font-weight: 700;
        display: flex; align-items: center; gap: 5px;
    }
    .rcm-status-dot {
        width: 6px; height: 6px; border-radius: 50%;
        background: #34C759;
        box-shadow: 0 0 6px #34C759;
    }
    .rcm-sync-btn-wrap { height: 100%; display: flex; align-items: center; }
    .rcm-sync-btn-wrap .stButton { width: 100%; }
    .rcm-sync-btn-wrap .stButton button {
        height: 100%;
        min-height: 66px;
        background: linear-gradient(180deg, #101E33, #0C1728) !important;
        border: 1px solid #1E3A5F !important;
        color: #C3D4EA !important;
        font-size: 12px !important;
        white-space: normal !important;
        line-height: 1.3;
    }
    .rcm-sync-btn-wrap .stButton button:hover {
        border-color: #FF8C42 !important;
        color: #F7FAFD !important;
    }

    /* Expander styled like clickable list card */
    div[data-testid="stExpander"] {
        border: 1px solid #1E3A5F !important;
        border-radius: 14px !important;
        background: linear-gradient(90deg, rgba(47,128,237,0.10), rgba(12,23,40,0.55)) !important;
        overflow: hidden;
    }
    div[data-testid="stExpander"] summary {
        padding: 14px 20px !important;
        font-weight: 600 !important;
    }

    /* ---------- Cards ---------- */
    .rcm-card {
        background: linear-gradient(180deg, #101E33 0%, #0C1728 100%);
        border: 1px solid #1E3A5F;
        border-radius: 14px;
        padding: 18px 22px;
        margin-bottom: 16px;
        box-shadow: 0 4px 18px rgba(0,0,0,0.22);
    }
    .rcm-card-success {
        border-color: #2A6E4C;
        background: linear-gradient(120deg, rgba(42,110,76,0.18), rgba(12,23,40,0.5));
    }
    .rcm-card-warn {
        border-color: #8A6329;
        background: linear-gradient(120deg, rgba(138,99,41,0.18), rgba(12,23,40,0.5));
    }

    /* ---------- Stat boxes ---------- */
    .rcm-stat {
        background: linear-gradient(180deg, #0F1D32 0%, #0A1526 100%);
        border: 1px solid #1E3A5F;
        border-radius: 12px;
        padding: 16px 18px;
        text-align: center;
        transition: border-color 0.18s ease, transform 0.18s ease, box-shadow 0.18s ease;
    }
    .rcm-stat:hover {
        border-color: #FF8C42;
        transform: translateY(-2px);
        box-shadow: 0 6px 20px rgba(255,107,53,0.15);
    }
    .rcm-stat-label {
        font-family: 'JetBrains Mono', monospace;
        font-size: 11px;
        color: #6B84A6;
        text-transform: uppercase;
        letter-spacing: 0.1em;
        margin-bottom: 8px;
    }
    .rcm-stat-value {
        font-size: 24px;
        font-weight: 800;
        color: #F7FAFD;
        background: linear-gradient(90deg, #F7FAFD, #B8D4F0);
        -webkit-background-clip: text;
        -webkit-text-fill-color: transparent;
    }
    .rcm-stat-sub {
        font-family: 'JetBrains Mono', monospace;
        font-size: 11px;
        color: #FF9B5C;
        margin-top: 4px;
    }

    /* ---------- Buttons ---------- */
    .stDownloadButton button, .stButton button {
        border-radius: 10px !important;
        font-weight: 700 !important;
        border: none !important;
        transition: transform 0.15s ease, box-shadow 0.15s ease !important;
    }
    .stDownloadButton button {
        background: linear-gradient(90deg, #FF8C42, #FF6B35) !important;
        color: #0B1220 !important;
        box-shadow: 0 4px 16px rgba(255,107,53,0.28) !important;
    }
    .stDownloadButton button:hover, .stButton button:hover {
        transform: translateY(-1px);
        box-shadow: 0 6px 22px rgba(255,107,53,0.4) !important;
    }
    button[kind="primary"] {
        background: linear-gradient(90deg, #2F80ED, #1E63C4) !important;
        color: #F7FAFD !important;
        box-shadow: 0 4px 16px rgba(47,128,237,0.3) !important;
    }

    /* Custom "preview" toggle button styled as a list card */
    .rcm-preview-toggle-wrap .stButton button {
        background: linear-gradient(90deg, rgba(47,128,237,0.10), rgba(12,23,40,0.55)) !important;
        border: 1px solid #1E3A5F !important;
        border-radius: 14px !important;
        color: #E8F0FA !important;
        text-align: left !important;
        justify-content: flex-start !important;
        padding: 16px 20px !important;
        font-weight: 500 !important;
        font-size: 13.5px !important;
        box-shadow: none !important;
    }
    .rcm-preview-toggle-wrap .stButton button:hover {
        border-color: #2F80ED !important;
        transform: none !important;
        box-shadow: 0 4px 14px rgba(47,128,237,0.18) !important;
    }

    [data-testid="stFileUploaderDropzone"] {
        background: linear-gradient(180deg, #0F1D32, #0A1526) !important;
        border: 2px dashed #2A4A72 !important;
        border-radius: 14px !important;
    }
    [data-testid="stFileUploaderDropzone"] button {
        background: linear-gradient(90deg, #2F80ED, #1E63C4) !important;
        border-radius: 8px !important;
        border: none !important;
        outline: none !important;
        box-shadow: none !important;
        position: relative !important;
        padding: 9px 20px !important;
        min-width: 150px;
        min-height: 38px;
    }
    [data-testid="stFileUploaderDropzone"] button:focus,
    [data-testid="stFileUploaderDropzone"] button:focus-visible {
        outline: none !important;
        box-shadow: none !important;
        border-color: transparent !important;
    }
    [data-testid="stFileUploaderDropzone"] button * {
        display: none !important;
    }
    [data-testid="stFileUploaderDropzone"] button::after {
        content: "↑  Upload File";
        position: absolute;
        inset: 0;
        display: flex !important;
        align-items: center;
        justify-content: center;
        font-size: 13.5px;
        font-weight: 600;
        color: #F7FAFD;
        font-family: 'Inter', sans-serif;
        white-space: nowrap;
    }
    [data-testid="stFileUploaderDropzoneInstructions"] span {
        font-size: 0 !important;
        display: block;
    }
    [data-testid="stFileUploaderDropzoneInstructions"] span::after {
        content: "ลากไฟล์ route มาวางที่นี่ หรือคลิกเพื่อเลือกไฟล์";
        font-size: 14px;
        font-weight: 600;
        color: #E8F0FA;
        font-family: 'Inter', sans-serif;
    }
    [data-testid="stFileUploaderDropzoneInstructions"] small {
        color: #6B84A6 !important;
        font-family: 'JetBrains Mono', monospace !important;
    }

    /* ---------- Dataframe ---------- */
    [data-testid="stDataFrame"] { border-radius: 12px; overflow: hidden; border: 1px solid #1E3A5F; }

    hr { border-color: #1E3A5F !important; }

    /* Radio pills for export format */
    div[role="radiogroup"] { gap: 8px; }

    /* subtle divider glow */
    .rcm-glow-divider {
        height: 1px;
        background: linear-gradient(90deg, transparent, #2A4A72, transparent);
        margin: 22px 0;
    }
    /* ---------- Region picker ---------- */
    [data-testid="stRadio"] > label p { color: #C3D4EA !important; font-weight: 600; font-size: 13px; letter-spacing: .3px; }
    [data-testid="stRadio"] [role="radiogroup"] label p { color: #F0F5FA !important; font-weight: 600; font-size: 15px; }
    [data-testid="stRadio"] [role="radiogroup"] { gap: 1.4rem; }

    /* ---------- Tabs: ตัวหนังสือสว่างตลอด ไม่ต้องกดก่อนถึงจะเห็น ---------- */
    [data-testid="stTab"], [role="tab"], button[data-baseweb="tab"] { color: #F0F5FA !important; }
    [data-testid="stTab"] p, [role="tab"] p, button[data-baseweb="tab"] p {
        color: #F0F5FA !important; font-weight: 600; font-size: 15px;
    }
    [data-testid="stTab"]:hover p, [role="tab"]:hover p, button[data-baseweb="tab"]:hover p { color: #FFB347 !important; }
    [data-testid="stTab"][aria-selected="true"] p, [role="tab"][aria-selected="true"] p,
    button[data-baseweb="tab"][aria-selected="true"] p { color: #FF8C42 !important; font-weight: 700; }
    [role="tablist"], [data-baseweb="tab-list"] { border-bottom: 1px solid rgba(195,212,234,0.18); }
    [data-baseweb="tab-highlight"], .react-aria-SelectionIndicator { background-color: #FF8C42 !important; }
</style>
""", unsafe_allow_html=True)

MASTER_SHEET_CSV_URL = st.secrets.get("MASTER_SHEET_CSV_URL", "")
PRODUCT_SHEET_CSV_URL = st.secrets.get("PRODUCT_SHEET_CSV_URL", "")


# ============================================================
# HELPERS
# ============================================================
def normalize(s):
    if s is None or (isinstance(s, float) and pd.isna(s)):
        return ''
    return re.sub(r'\s+', ' ', str(s).strip().upper().replace('.', '').replace(',', ''))


@st.cache_data(ttl=300)
def load_master_data(url):
    df = pd.read_csv(url)
    df.columns = [c.strip() for c in df.columns]
    required = ['Ship To', 'Ship To Name', 'Latitude', 'Longitude']
    missing = [c for c in required if c not in df.columns]
    if missing:
        raise ValueError(f"ไม่พบคอลัมน์: {', '.join(missing)} กรุณาตรวจสอบหัวตารางใน Google Sheets")
    return prepare_master(df)


def prepare_master(df):
    """Customer Master -> แถวที่ใช้จับคู่พิกัด
    - ต้องมี Latitude + Longitude (ไม่มีพิกัดก็ใช้ไม่ได้)
    - แถวที่ไม่มีรหัส Ship To ยังใช้จับคู่ด้วยชื่อได้ แต่เฉพาะชื่อที่ยังไม่มีแถวที่มีรหัส
      (ถ้าชื่อนี้มีแถวที่มีรหัสอยู่แล้ว ใช้แถวที่มีรหัสเสมอ — แถวที่เพิ่มเองไม่ไปทับข้อมูลจากระบบ)"""
    df = df.dropna(subset=['Latitude', 'Longitude']).copy()
    df = df[df['Ship To Name'].notna() | df['Ship To'].notna()]
    df['norm_ship'] = df['Ship To Name'].apply(normalize)
    has_code = df['Ship To'].map(lambda v: not _cell_blank(v))
    df['norm_code'] = [str(v).strip().upper() if h else None for v, h in zip(df['Ship To'], has_code)]
    coded_names = set(df.loc[has_code, 'norm_ship'])
    # กันแถวที่เพิ่มเองเป็น "ชื่อบริษัทแม่" (เช่น BAKEHOUSE CO.,LTD.) ที่บริษัทนั้นมีหลายสาขาคนละที่ในระบบ
    # ถ้าใช้ จะได้พิกัดสำนักงานใหญ่ไปใส่ให้ทุกสาขา -> ไม่ใช้แถวแบบนี้
    if 'Name' in df.columns:
        coded = df[has_code].assign(_cn=df.loc[has_code, 'Name'].apply(normalize))
        pins = coded.groupby('_cn')[['Latitude', 'Longitude']].nunique()
        multi_branch_companies = set(pins[(pins['Latitude'] > 1) | (pins['Longitude'] > 1)].index)
    else:
        multi_branch_companies = set()
    keep_nocode = ~has_code & ~df['norm_ship'].isin(coded_names) & ~df['norm_ship'].isin(multi_branch_companies)
    df = df[has_code | keep_nocode]
    return df


def master_multi_pin_names(master_df):
    """ชื่อ (normalize แล้ว) ที่มีหลายแถวใน Master และพิกัดไม่ตรงกัน"""
    grp = master_df.groupby('norm_ship')[['Latitude', 'Longitude']].nunique()
    return set(grp[(grp['Latitude'] > 1) | (grp['Longitude'] > 1)].index)


@st.cache_data(ttl=300)
def load_product_master(url):
    # Product Master มี 2 แถวหัวเรื่อง (Ownership, แผนก) อยู่เหนือแถวชื่อคอลัมน์จริง
    # แถวชื่อคอลัมน์จริงคือแถวที่ 3 (index 2)
    df = pd.read_csv(url, header=2)
    df.columns = [c.strip() for c in df.columns]
    required = ['SKUCode', 'PRODUCT']
    missing = [c for c in required if c not in df.columns]
    if missing:
        raise ValueError(f"ไม่พบคอลัมน์: {', '.join(missing)} กรุณาตรวจสอบหัวตารางใน Google Sheets")
    df = df.dropna(subset=['SKUCode'])
    df['SKUCode'] = df['SKUCode'].astype(str).str.strip()
    return df


def find_header_row(rows, keys):
    for r in range(min(len(rows), 10)):
        row = rows[r]
        for c, v in enumerate(row):
            if isinstance(v, str) and v.strip().lower() in keys:
                return r, c
    return None, None


def _cell_blank(v):
    return v is None or (isinstance(v, float) and pd.isna(v)) or str(v).strip() == ''


def _fill_merged_down(raw, ws, header_row):
    """ไฟล์บางแบบ (เช่น Route Planning สมุย) ใช้ merge cell แนวตั้ง: จุดส่ง 1 จุดมีหลาย PO
    แต่ Cust. ID / ชื่อลูกค้า / Ship-to ถูก merge ไว้แค่แถวบนสุด pandas จะเห็นแถวล่างๆ เป็นช่องว่าง
    -> แถว PO พวกนั้นจะไม่มี Cust ID และหลุดหายจากผลลัพธ์ ฟังก์ชันนี้เติมค่าจาก merge ลงไปทุกแถว
    ในช่วงนั้น (ทำเฉพาะ merge แนวตั้งที่อยู่ใต้หัวตาราง และเฉพาะเมื่อค่ามุมบนซ้ายตรงกันระหว่าง
    pandas กับ openpyxl — ถ้าไม่ตรงแปลว่าแถวไม่ตรงกัน ข้ามไปเลย ไม่เดา)"""
    for rng in ws.merged_cells.ranges:
        r0, c0 = rng.min_row - 1, rng.min_col - 1
        r1, c1 = rng.max_row - 1, rng.max_col - 1
        if r0 <= header_row or c0 != c1 or r1 >= len(raw):
            continue
        top = raw[r0][c0] if c0 < len(raw[r0]) else None
        anchor = ws.cell(row=rng.min_row, column=rng.min_col).value
        if _cell_blank(top) or anchor is None or str(top).strip() != str(anchor).strip():
            continue
        for r in range(r0 + 1, r1 + 1):
            if c0 < len(raw[r]) and _cell_blank(raw[r][c0]):
                raw[r][c0] = top


def process_route_file(file_bytes, master_df, selected_sheets=None, dedupe_shipto=True):
    xl = pd.ExcelFile(io.BytesIO(file_bytes))
    # ใช้อ่าน merge cell (.xlsx เท่านั้น — .xls เปิดด้วย openpyxl ไม่ได้ ก็ข้ามไปตามปกติ)
    try:
        merge_wb = load_workbook(io.BytesIO(file_bytes))
    except Exception:
        merge_wb = None
    output_sheets = {}
    unique_stops_sheets = {}  # เวอร์ชัน "จุดส่งไม่ซ้ำ" คอลัมน์เหมือนไฟล์หลักทุกอย่าง แค่ยุบแถวพิกัดซ้ำ
    report = []
    unmatched = []
    ambiguous = []

    by_ship = master_df.drop_duplicates('norm_ship').set_index('norm_ship')
    by_code = master_df.dropna(subset=['norm_code']).drop_duplicates('norm_code').set_index('norm_code')
    multi_pin_keys = master_multi_pin_names(master_df)

    ambiguous_ship_keys = set()
    if not dedupe_shipto:
        # หา Ship To Name ที่ซ้ำกันแต่พิกัด "ขัดแย้งกันจริง" (คนละสาขา) เพื่อไม่เดาสุ่มเลือกสาขาใดสาขาหนึ่ง
        grp = master_df.groupby('norm_ship')[['Latitude', 'Longitude']].nunique()
        ambiguous_ship_keys = set(grp[(grp['Latitude'] > 1) | (grp['Longitude'] > 1)].index)

    sheet_list = selected_sheets if selected_sheets is not None else xl.sheet_names

    for sheet_name in sheet_list:
        raw = xl.parse(sheet_name, header=None).values.tolist()
        header_row, cust_col = find_header_row(raw, [
            'cust code', 'cust id', 'cust. id', 'cust.id', 'cust. code', 'cust.code',
            'customer code', 'customers code', 'customer id', 'customers id',
        ])

        # ถ้าไม่เจอ Cust Code/ID เลย ลองหาคอลัมน์ทางเลือก (ไฟล์บางแบบ เช่น export สไตล์ invoice
        # ของโซนภูเก็ต/สมุย ไม่มีคอลัมน์ Cust Code เลย มีแต่ Cust. Name)
        if header_row is None:
            header_row, cust_col = find_header_row(raw, ['cust. name', 'cust name', 'cust.name', 'customer name', 'customers name'])

        if header_row is None:
            output_sheets[sheet_name] = xl.parse(sheet_name, header=None)
            report.append({'sheet': sheet_name, 'skipped': True, 'matched': 0, 'total': 0})
            continue

        if merge_wb is not None and sheet_name in merge_wb.sheetnames:
            _fill_merged_down(raw, merge_wb[sheet_name], header_row)

        header = raw[header_row]
        ship_col = None
        for i, v in enumerate(header):
            if isinstance(v, str) and 'ship' in v.lower() and 'name' in v.lower():
                ship_col = i
                break
        # ถ้าไม่เจอคอลัมน์ชื่อ "...ship...name..." ลองหา "Ship To" เฉยๆ (ไม่มีคำว่า name ต่อท้าย)
        # เผื่อไฟล์ตั้งชื่อคอลัมน์แบบนี้ (เจอในไฟล์ export บางโซน)
        if ship_col is None:
            for i, v in enumerate(header):
                if isinstance(v, str) and v.strip().lower() in ('ship to', 'ship-to', 'shipto'):
                    ship_col = i
                    break

        # คอลัมน์ Ship To มีอยู่ แต่ว่างทั้งชีท (เช่นชีท 23.09.26 ของไฟล์สมุย) = เท่ากับไม่มีคอลัมน์นี้
        # -> ตัดสินใจระดับชีทให้ไปใช้ชื่อลูกค้าแทน (ยังเป็นการจับด้วยชื่อ ไม่ใช่เดาด้วยรหัส)
        if ship_col is not None and all(
            ship_col >= len(rw) or _cell_blank(rw[ship_col]) for rw in raw[header_row + 1:]
        ):
            ship_col = None

        # ถ้าไม่มีคอลัมน์ ShipTo Name ในชีทนี้เลย ให้ใช้ Cust Name (หรือคอลัมน์ชื่ออื่นที่ใกล้เคียง) แทน
        name_fallback_col = None
        if ship_col is None:
            for i, v in enumerate(header):
                if isinstance(v, str) and 'cust' in v.lower() and 'name' in v.lower():
                    name_fallback_col = i
                    break
            if name_fallback_col is None:
                for i, v in enumerate(header):
                    if isinstance(v, str) and 'name' in v.lower() and 'ship' not in v.lower():
                        name_fallback_col = i
                        break

        # ระบุ "โหมด" การจับคู่ของทั้งชีทนี้ ครั้งเดียวตอนเจอ header (ไม่ใช่ทีละแถว):
        #   - ถ้ามีคอลัมน์ ShipTo Name หรือ Cust Name -> ใช้ชื่อจับคู่กับ Master ทุกแถวในชีทนี้
        #     (แถวไหน match ไม่ได้ ให้ถือว่า "ไม่มีพิกัด" ไปเลย ห้ามหลุดไปเดาด้วย Cust Code
        #      เพราะลูกค้า 1 รายอาจมีหลายสาขา การเดาด้วยโค้ดระดับลูกค้าจะได้พิกัดสาขาอื่นมาแบบผิดๆ)
        #   - ถ้าชีทนี้ไม่มีทั้งสองคอลัมน์เลย -> ทั้งชีทจะจับคู่ด้วย Cust Code <-> Ship To (Master) แทน
        use_code_fallback_for_sheet = (ship_col is None and name_fallback_col is None)
        # มีคอลัมน์ ShipTo Name แต่บางแถวว่าง (เช่นแถวที่ทีมเพิ่มเองใน Sheet2) -> แถวนั้นใช้ Customer Name แทน
        row_name_col = None
        if ship_col is not None:
            for i, v in enumerate(header):
                if isinstance(v, str) and ('cust' in v.lower() or 'customer' in v.lower()) and 'name' in v.lower():
                    row_name_col = i
                    break

        out_rows = []
        extra_rows = set()  # index ใน out_rows ของแถวที่ไม่ใช่ออเดอร์ (เก็บไว้แต่ไม่เอาไปไฟล์จุดส่งไม่ซ้ำ)
        matched = total = via_ship = via_code = 0

        for r, row in enumerate(raw):
            if r < header_row:
                out_rows.append(row)
                continue
            if r == header_row:
                new_header = list(row) + ['Latitude', 'Longitude']
                out_rows.append(new_header)
                continue

            cust_val = row[cust_col] if cust_col < len(row) else None
            if cust_val is None or (isinstance(cust_val, float) and pd.isna(cust_val)) or str(cust_val).strip() == '':
                # ไม่มีรหัสลูกค้า: ถ้าแถวมีข้อความ (เช่น ท้ายใบงาน "ลงชื่อ...พนักงานขับ", หมายเหตุถึงคนขับ,
                # งาน "ตามใบคุม") -> เก็บไว้ตามเดิม ไม่มีพิกัด และไม่นับเป็นออเดอร์
                # ตัดทิ้งเฉพาะแถวว่าง หรือแถวที่มีแค่ช่องแรก (เช่นเลขลำดับเปล่าๆ ในเทมเพลตสมุย)
                if any(not _cell_blank(v) for v in row[1:]):
                    out_rows.append(list(row) + [None, None])
                    extra_rows.add(len(out_rows) - 1)
                continue

            code = str(cust_val).strip()
            total += 1
            via_row_fallback = False
            if ship_col is not None and ship_col < len(row) and not _cell_blank(row[ship_col]):
                ship_val = row[ship_col]
            elif ship_col is not None:
                # ShipTo Name ว่างเฉพาะแถวนี้ -> ใช้ Customer Name (ถ้าชื่อนี้มีหลายหมุดใน Master จะไม่เดา)
                ship_val = row[row_name_col] if (row_name_col is not None and row_name_col < len(row)
                                                 and not _cell_blank(row[row_name_col])) else None
                via_row_fallback = ship_val is not None
            elif name_fallback_col is not None and name_fallback_col < len(row):
                ship_val = row[name_fallback_col]
            else:
                ship_val = None

            hit = None
            is_ambiguous = False
            if not use_code_fallback_for_sheet:
                # โหมดจับคู่ด้วยชื่อ: ลองครั้งเดียว ไม่ fallback ไป Cust Code ต่อแถวต่อแถว
                if ship_val:
                    key = normalize(ship_val)
                    if key in ambiguous_ship_keys or (via_row_fallback and key in multi_pin_keys):
                        # ชื่อนี้ซ้ำในหลายสาขาที่พิกัดขัดแย้งกัน (เช่น Big C หลายสาขาชื่อเดียวกัน)
                        # ไม่เดาว่าเป็นสาขาไหน -> ถือว่ายังไม่มีพิกัด แล้วบันทึกแยกไว้ให้ตรวจสอบ
                        is_ambiguous = True
                    elif key in by_ship.index:
                        hit = by_ship.loc[key]
                        via_ship += 1
            else:
                # โหมดนี้ใช้เมื่อทั้งชีทไม่มีคอลัมน์ชื่อเลย -> จับคู่ด้วย Cust Code <-> Ship To (Master)
                key = code.upper()
                if key in by_code.index:
                    hit = by_code.loc[key]
                    via_code += 1

            new_row = list(row)
            if hit is not None:
                matched += 1
                lat = hit['Latitude'].iloc[0] if hasattr(hit['Latitude'], 'iloc') else hit['Latitude']
                lon = hit['Longitude'].iloc[0] if hasattr(hit['Longitude'], 'iloc') else hit['Longitude']
                new_row += [lat, lon]
            elif is_ambiguous:
                new_row += [None, None]
                dup_rows = master_df[master_df['norm_ship'] == normalize(ship_val)]
                branch_count = len(dup_rows.drop_duplicates(subset=['Latitude', 'Longitude']))
                ambiguous.append({
                    'Cust Code': code,
                    'Ship To Name': ship_val or '',
                    'Sheet': sheet_name,
                    'จำนวนสาขาที่ชื่อซ้ำ': branch_count,
                    'พบเมื่อ': datetime.now().strftime('%Y-%m-%d'),
                })
            else:
                new_row += [None, None]
                unmatched.append({
                    'Cust Code': code,
                    'Ship To Name': ship_val or '',
                    'Sheet': sheet_name,
                    'พบเมื่อ': datetime.now().strftime('%Y-%m-%d'),

                })
            out_rows.append(new_row)

        output_sheets[sheet_name] = pd.DataFrame(out_rows)
        report.append({'sheet': sheet_name, 'skipped': False, 'matched': matched, 'total': total,
                        'via_ship': via_ship, 'via_code': via_code})

        # ---- สร้างเวอร์ชัน "จุดส่งไม่ซ้ำ" ของชีทนี้ ----
        # ใช้ template เดียวกับ out_rows เป๊ะ (คอลัมน์เดิมทั้งหมด + Latitude/Longitude)
        # แค่ยุบแถวที่พิกัด (2 คอลัมน์สุดท้าย) ซ้ำกันให้เหลือแถวแรกที่เจอ แล้วเติมคอลัมน์นับจำนวนออเดอร์ต่อท้าย
        preamble_and_header = out_rows[:header_row + 1]
        data_rows = [r_ for k_, r_ in enumerate(out_rows) if k_ > header_row and k_ not in extra_rows]

        seen_order = []
        seen_row = {}
        seen_count = {}
        no_coord_counter = 0
        for row in data_rows:
            lat_val, lon_val = row[-2], row[-1]
            if lat_val is None or lon_val is None:
                # ไม่มีพิกัด -> ยังคงเก็บแถวนี้ไว้เหมือนเดิม (ห้ามลบทิ้งเงียบๆ) แค่ไม่เอาไปยุบรวมกับแถวไหน
                # เพราะไม่รู้ว่ามันคือจุดเดียวกับแถวอื่นไหนหรือเปล่า
                key = ('__NOCOORD__', no_coord_counter)
                no_coord_counter += 1
            else:
                key = (lat_val, lon_val)
            if key not in seen_row:
                seen_row[key] = row
                seen_count[key] = 1
                seen_order.append(key)
            else:
                seen_count[key] += 1

        unique_header = list(preamble_and_header[-1]) + ['จำนวนออเดอร์รวม'] if preamble_and_header else \
            list(raw[header_row]) + ['Latitude', 'Longitude', 'จำนวนออเดอร์รวม']
        unique_data_rows = [list(seen_row[k]) + [seen_count[k]] for k in seen_order]
        unique_sheet_rows = preamble_and_header[:-1] + [unique_header] + unique_data_rows if preamble_and_header else \
            [unique_header] + unique_data_rows
        unique_stops_sheets[sheet_name] = pd.DataFrame(unique_sheet_rows)
        report[-1]['unique_stops'] = len(unique_data_rows)

    return output_sheets, report, unmatched, ambiguous, unique_stops_sheets


def process_weight_file(file_bytes, product_df, selected_sheets=None):
    """เชื่อม Product Code (ไฟล์ route) กับ SKUCode (Product Master) เพื่อคำนวณน้ำหนักทีละแถว
    เก็บคอลัมน์เดิมทั้งหมดไว้ครบ (เหมือนโหมดพิกัด) แค่เพิ่ม 'Weight (kg)' ต่อท้าย"""
    xl = pd.ExcelFile(io.BytesIO(file_bytes))
    output_sheets = {}
    report = []
    no_weight = []

    weight_lookup = product_df.dropna(subset=['SKUCode']).drop_duplicates('SKUCode').set_index('SKUCode')

    sheet_list = selected_sheets if selected_sheets is not None else xl.sheet_names

    for sheet_name in sheet_list:
        raw = xl.parse(sheet_name, header=None).values.tolist()
        header_row, product_col = find_header_row(raw, ['product code'])

        if header_row is None:
            output_sheets[sheet_name] = xl.parse(sheet_name, header=None)
            report.append({'sheet': sheet_name, 'skipped': True, 'calculated': 0, 'total': 0})
            continue

        header = raw[header_row]
        qty_col = None
        units_col = None
        for i, v in enumerate(header):
            if isinstance(v, str) and v.strip().lower() == 'qty':
                qty_col = i
            if isinstance(v, str) and v.strip().lower() == 'units':
                units_col = i

        out_rows = []
        calculated = total = 0

        for r, row in enumerate(raw):
            if r < header_row:
                out_rows.append(row)
                continue
            if r == header_row:
                out_rows.append(list(row) + ['Weight (kg)'])
                continue

            product_val = row[product_col] if product_col < len(row) else None
            if product_val is None or (isinstance(product_val, float) and pd.isna(product_val)) or str(product_val).strip() == '':
                # ไม่มี Product Code -> ยังคงเก็บแถวนี้ไว้เหมือนเดิม (ห้ามลบทิ้งเงียบๆ) แค่ไม่มีน้ำหนักให้คำนวณ
                out_rows.append(list(row) + [None])
                continue

            code = str(product_val).strip()
            total += 1
            qty_val = row[qty_col] if (qty_col is not None and qty_col < len(row)) else None
            units_val = row[units_col] if (units_col is not None and units_col < len(row)) else None
            units_str = str(units_val).strip().upper() if units_val is not None and not (isinstance(units_val, float) and pd.isna(units_val)) else ''

            weight = None
            reason = None

            if code not in weight_lookup.index:
                reason = 'ไม่พบ Product Code นี้ใน Product Master'
            elif qty_val is None or (isinstance(qty_val, float) and pd.isna(qty_val)):
                reason = 'ไม่มีค่า Qty ในแถวนี้'
            elif units_str == 'PCS':
                unit_wt = weight_lookup.loc[code, 'Unit Net Wt (kg)']
                if hasattr(unit_wt, 'iloc'):
                    unit_wt = unit_wt.iloc[0]
                if pd.isna(unit_wt):
                    reason = 'มี Product Code แต่ Product Master ยังไม่มีข้อมูลน้ำหนัก (Unit Net Wt)'
                else:
                    weight = float(qty_val) * float(unit_wt)
            elif units_str == 'KG':
                weight = float(qty_val)  # Qty เป็นน้ำหนักอยู่แล้ว ไม่ต้องคูณ
            else:
                reason = f'หน่วย (Units) "{units_val}" ยังไม่รองรับการคำนวณอัตโนมัติ'

            new_row = list(row) + [weight]
            if weight is not None:
                calculated += 1
            else:
                no_weight.append({
                    'Product Code': code,
                    'Qty': qty_val,
                    'Units': units_val if units_val is not None else '',
                    'Sheet': sheet_name,
                    'สาเหตุ': reason,
                })
            out_rows.append(new_row)

        output_sheets[sheet_name] = pd.DataFrame(out_rows)
        report.append({'sheet': sheet_name, 'skipped': False, 'calculated': calculated, 'total': total})

    return output_sheets, report, no_weight


def _is_blank(v):
    if v is None:
        return True
    if isinstance(v, float) and pd.isna(v):
        return True
    if isinstance(v, str) and v.strip() == '':
        return True
    return False


def _drop_empty_columns(df):
    """ลบคอลัมน์ที่ว่างเปล่าทุกแถว (ทั้ง header และข้อมูล) เช่นคอลัมน์ spacer ระหว่างคอลัมน์จริงในไฟล์ route ต้นฉบับ"""
    keep_cols = [c for c in df.columns if not df[c].apply(_is_blank).all()]
    if not keep_cols:  # กันเคส edge case ที่ทุกคอลัมน์ว่างหมด (ไม่ควรเกิดขึ้นจริง)
        return df
    return df[keep_cols].reset_index(drop=True)


def to_excel_bytes(sheets_dict):
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        for name, df in sheets_dict.items():
            df_clean = _drop_empty_columns(df)
            df_clean.to_excel(writer, sheet_name=name[:31], header=False, index=False)
            ws = writer.sheets[name[:31]]
            # ปรับความกว้างคอลัมน์ให้พอดีกับตัวอักษรที่ยาวที่สุดในคอลัมน์นั้น (openpyxl ไม่มี autofit ในตัว จำลองเอง)
            for col_idx in range(df_clean.shape[1]):
                max_len = 0
                for val in df_clean.iloc[:, col_idx]:
                    if not _is_blank(val):
                        max_len = max(max_len, len(str(val)))
                col_letter = get_column_letter(col_idx + 1)
                ws.column_dimensions[col_letter].width = min(max(max_len + 2, 8), 60)
    return output.getvalue()


def to_csv_bytes(sheets_dict):
    # รวมทุก sheet ที่เลือกเป็น CSV เดียว คั่นด้วยชื่อ sheet
    output = io.StringIO()
    for name, df in sheets_dict.items():
        output.write(f"# Sheet: {name}\n")
        df.to_csv(output, header=False, index=False)
        output.write("\n")
    return output.getvalue().encode('utf-8-sig')


def to_kml_bytes(sheets_dict):
    # หาแถว header จริง (แถวที่มีคำว่า Latitude/Longitude) เพราะบาง sheet มีแถวหัวเรื่องอยู่ด้านบน header จริง
    placemarks = []
    for sheet_name, df in sheets_dict.items():
        if df.empty or df.shape[1] < 2:
            continue

        header_row_idx = None
        lat_idx = lon_idx = None
        for i in range(min(len(df), 15)):
            row_vals = df.iloc[i].tolist()
            if 'Latitude' in row_vals and 'Longitude' in row_vals:
                header_row_idx = i
                lat_idx = row_vals.index('Latitude')
                lon_idx = row_vals.index('Longitude')
                break
        if header_row_idx is None:
            continue

        header = df.iloc[header_row_idx].tolist()
        name_idx = 0
        for i, h in enumerate(header):
            if isinstance(h, str) and 'name' in h.lower():
                name_idx = i
                break

        for _, row in df.iloc[header_row_idx + 1:].iterrows():
            lat, lon = row.iloc[lat_idx], row.iloc[lon_idx]
            if pd.isna(lat) or pd.isna(lon):
                continue
            label = str(row.iloc[name_idx]) if name_idx < len(row) else "Stop"
            placemarks.append(f"""
    <Placemark>
      <name>{label}</name>
      <description>Sheet: {sheet_name}</description>
      <Point><coordinates>{lon},{lat},0</coordinates></Point>
    </Placemark>""")

    kml = f"""<?xml version="1.0" encoding="UTF-8"?>
<kml xmlns="http://www.opengis.net/kml/2.2">
  <Document>
    <name>Route Coordinates</name>
    {''.join(placemarks)}
  </Document>
</kml>"""
    return kml.encode('utf-8')


def unmatched_to_excel_bytes(unmatched_df):
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        unmatched_df.to_excel(writer, sheet_name='Unmatched', index=False)
        ws = writer.sheets['Unmatched']
        for i, col in enumerate(unmatched_df.columns, start=1):
            ws.column_dimensions[chr(64 + i)].width = 28
    return output.getvalue()


def stat_box(label, value, sub=""):
    st.markdown(f"""
    <div class="rcm-stat">
        <div class="rcm-stat-label">{label}</div>
        <div class="rcm-stat-value">{value}</div>
        <div class="rcm-stat-sub">{sub}</div>
    </div>
    """, unsafe_allow_html=True)


# ============================================================
# HEADER
# ============================================================
st.markdown("""
<div class="rcm-header">
    <svg class="rcm-route-svg" viewBox="0 0 400 220" xmlns="http://www.w3.org/2000/svg">
        <path d="M 60 190 Q 150 170 180 120 T 300 40" stroke="rgba(255,255,255,0.22)" stroke-width="1.5" fill="none" stroke-dasharray="4 5"/>
        <path d="M 90 200 Q 170 150 220 130 T 330 55" stroke="rgba(255,255,255,0.12)" stroke-width="1" fill="none"/>
        <circle cx="300" cy="40" r="7" fill="#FF6B35" opacity="0.95">
            <animate attributeName="opacity" values="0.6;1;0.6" dur="2.4s" repeatCount="indefinite"/>
        </circle>
        <circle cx="300" cy="40" r="13" fill="none" stroke="#FF6B35" stroke-width="1.5" opacity="0.5"/>
        <circle cx="60" cy="190" r="6" fill="#2F80ED" opacity="0.95">
            <animate attributeName="opacity" values="0.6;1;0.6" dur="2.4s" repeatCount="indefinite" begin="1.2s"/>
        </circle>
        <circle cx="60" cy="190" r="11" fill="none" stroke="#2F80ED" stroke-width="1.5" opacity="0.5"/>
        <circle cx="180" cy="120" r="3.5" fill="#FFB347" opacity="0.8"/>
    </svg>
    <div class="rcm-header-content">
        <div class="rcm-eyebrow">🚚 FLEET OPERATIONS · GEO-MATCHING SYSTEM</div>
        <div class="rcm-title">Route Coordinate Matcher</div>
        <div class="rcm-subtitle">อัปโหลดไฟล์ route เพื่อจับคู่พิกัดกับ Master Data</div>
        <div class="rcm-features">
            <div class="rcm-feature">
                <div class="rcm-feature-icon">🎯</div>
                <div class="rcm-feature-text">
                    <div class="rcm-feature-title">Geo-Matching</div>
                    <div class="rcm-feature-sub">จับคู่พิกัดอัตโนมัติ</div>
                </div>
            </div>
            <div class="rcm-feature">
                <div class="rcm-feature-icon">🗄️</div>
                <div class="rcm-feature-text">
                    <div class="rcm-feature-title">Master Data</div>
                    <div class="rcm-feature-sub">ฐานข้อมูลลูกค้า</div>
                </div>
            </div>
            <div class="rcm-feature">
                <div class="rcm-feature-icon">🚛</div>
                <div class="rcm-feature-text">
                    <div class="rcm-feature-title">Fleet Efficiency</div>
                    <div class="rcm-feature-sub">เพิ่มประสิทธิภาพการขนส่ง</div>
                </div>
            </div>
        </div>
    </div>
</div>
""", unsafe_allow_html=True)

# ============================================================
# เลือกพื้นที่ (จังหวัด) — แต่ละพื้นที่เห็นเฉพาะเครื่องมือที่ใช้จริง (ดู REGION_TABS ท้ายไฟล์)
# ============================================================
REGION_LABELS = {'BKK': '🏙️ BKK · กรุงเทพ', 'PKT': '🏝️ PKT · ภูเก็ต', 'SAMUI': '🌴 SAMUI · สมุย'}
region = st.radio(
    "📍 เลือกพื้นที่", options=list(REGION_LABELS), format_func=REGION_LABELS.get,
    horizontal=True, key="region",
)

def _render_coords():
    if not MASTER_SHEET_CSV_URL:
        st.error("⚠️ ยังไม่ได้ตั้งค่า MASTER_SHEET_CSV_URL — ดูวิธีตั้งค่าใน DEPLOY_INSTRUCTIONS.md")
        return

    try:
        master_df = load_master_data(MASTER_SHEET_CSV_URL)
        unique_customer_count = master_df['norm_ship'].nunique()
        status_col, sync_col = st.columns([5, 1])
        with status_col:
            st.markdown(f"""
            <div class="rcm-status-card">
                <div class="rcm-status-left">
                    <div class="rcm-status-icon">✅</div>
                    <div>
                        <div class="rcm-status-title">เชื่อมต่อฐานข้อมูลสำเร็จ</div>
                        <div class="rcm-status-sub">{unique_customer_count:,} รายชื่อลูกค้าในระบบ (ไม่ซ้ำ)</div>
                    </div>
                </div>
                <div class="rcm-status-right">
                    <div class="rcm-status-right-icon">🗃️</div>
                    <div class="rcm-status-pill"><span class="rcm-status-dot"></span>Connected</div>
                </div>
            </div>
            """, unsafe_allow_html=True)
        with sync_col:
            st.markdown('<div class="rcm-sync-btn-wrap">', unsafe_allow_html=True)
            if st.button("🔄 ซิงค์ตอนนี้", use_container_width=True, help="ดึงข้อมูล Master Data ล่าสุดจาก Google Sheets ทันที"):
                load_master_data.clear()
                st.rerun()
            st.markdown('</div>', unsafe_allow_html=True)
    except Exception as e:
        st.error(f"โหลด Master Data ไม่สำเร็จ: {e}")
        return

    if 'show_master_preview' not in st.session_state:
        st.session_state['show_master_preview'] = False

    chevron = "▾" if st.session_state['show_master_preview'] else "▸"
    st.markdown('<div class="rcm-preview-toggle-wrap">', unsafe_allow_html=True)
    preview_clicked = st.button(
        f"📄   ดูตัวอย่าง Master Data — ตรวจสอบข้อมูลลูกค้าและเส้นทางก่อนประมวลผล   {chevron}",
        key="master_preview_toggle",
        use_container_width=True,
    )
    st.markdown('</div>', unsafe_allow_html=True)
    if preview_clicked:
        st.session_state['show_master_preview'] = not st.session_state['show_master_preview']
        st.rerun()

    if st.session_state['show_master_preview']:
        st.dataframe(master_df[['Ship To', 'Ship To Name', 'Latitude', 'Longitude']].head(20), use_container_width=True)

    st.write("")

    uploaded_file = st.file_uploader("อัพโหลดไฟล์ route", type=['xls', 'xlsx'], label_visibility="collapsed")

    if uploaded_file is not None:
        file_bytes = uploaded_file.read()

        try:
            xl_peek = pd.ExcelFile(io.BytesIO(file_bytes))
            all_sheet_names = xl_peek.sheet_names
        except Exception as e:
            st.error(f"เปิดไฟล์ไม่ได้: {e}")
            return

        st.markdown("##### 📑 เลือก Sheet ที่ต้องการประมวลผล")
        sheet_cols = st.columns(min(len(all_sheet_names), 4))
        selected_sheets = []
        for i, sn in enumerate(all_sheet_names):
            with sheet_cols[i % len(sheet_cols)]:
                checked = st.checkbox(sn, value=True, key=f"sheet_{sn}")
                if checked:
                    selected_sheets.append(sn)

        if not selected_sheets:
            st.warning("⚠️ กรุณาเลือกอย่างน้อย 1 Sheet เพื่อเริ่มประมวลผล")
            return

        dedupe_shipto = st.checkbox(
            "✨ ทำให้ Ship To Name ไม่ซ้ำกันอัตโนมัติ (เลือกแถวแรกที่เจอ)",
            value=True,
            help="เปิด (ค่าเริ่มต้น): ถ้าชื่อซ้ำกันในระบบ จะเลือกแถวแรกที่เจอมาใช้ทันที เร็ว แต่ถ้าเป็นเชนร้านที่มีหลายสาขาชื่อเดียวกัน (เช่น Big C) อาจได้พิกัดผิดสาขา\n\nปิด: ถ้าเจอชื่อซ้ำที่พิกัดขัดแย้งกันจริง (คนละสาขา) จะไม่เดา — ถือว่ายังไม่มีพิกัด แล้วแยกไปอยู่ลิสต์ 'ชื่อซ้ำ-ต้องตรวจสอบ' แทน ปลอดภัยกว่าแต่ match ได้น้อยลง",
        )

        st.write("")
        process_clicked = st.button("🚀 ประมวลผล Sheet ที่เลือก", type="primary", use_container_width=True)

        if process_clicked:
            with st.spinner("🛰️ กำลังจับคู่พิกัด..."):
                try:
                    sheets, report, unmatched, ambiguous, unique_stops_sheets = process_route_file(
                        file_bytes, master_df, selected_sheets, dedupe_shipto=dedupe_shipto
                    )
                    st.session_state['last_sheets'] = sheets
                    st.session_state['last_report'] = report
                    st.session_state['last_unmatched'] = unmatched
                    st.session_state['last_ambiguous'] = ambiguous
                    st.session_state['last_unique_stops_sheets'] = unique_stops_sheets
                    st.session_state['last_filename'] = uploaded_file.name
                except Exception as e:
                    st.error(f"เกิดข้อผิดพลาด: {e}")
                    st.exception(e)

        if 'last_sheets' in st.session_state and st.session_state.get('last_filename') == uploaded_file.name:
            sheets = st.session_state['last_sheets']
            report = st.session_state['last_report']
            unmatched = st.session_state['last_unmatched']
            ambiguous = st.session_state.get('last_ambiguous', [])
            unique_stops_sheets = st.session_state.get('last_unique_stops_sheets', {})

            total_matched = sum(r['matched'] for r in report)
            total_rows = sum(r['total'] for r in report)
            match_rate = (total_matched / total_rows * 100) if total_rows else 0

            card_class = "rcm-card-success" if match_rate >= 90 else "rcm-card-warn"
            st.markdown(f"""
            <div class="rcm-card {card_class}">
                🎯 <b>เสร็จแล้ว!</b> จับคู่พิกัดได้ {total_matched:,} / {total_rows:,} แถว ({match_rate:.1f}%)
            </div>
            """, unsafe_allow_html=True)

            cols = st.columns(len(report))
            for col, r in zip(cols, report):
                with col:
                    if r['skipped']:
                        stat_box(f"SHEET: {r['sheet']}", "ข้าม", "ไม่พบคอลัมน์ Cust Code")
                    else:
                        stat_box(f"SHEET: {r['sheet']}", f"{r['matched']}/{r['total']}",
                                  f"SHIP-TO {r['via_ship']} · CODE {r['via_code']}")

            st.markdown('<div class="rcm-glow-divider"></div>', unsafe_allow_html=True)

            # --- Export format selection ---
            st.markdown("##### 💾 เลือกฟอร์แมตไฟล์ผลลัพธ์")
            export_format = st.radio(
                "รูปแบบไฟล์",
                options=["Excel (.xlsx)", "CSV (.csv)", "KML (.kml — เปิดใน Google Earth/Maps)"],
                horizontal=True,
                label_visibility="collapsed",
            )

            base_name = uploaded_file.name.rsplit('.', 1)[0]
            dl_col1, dl_col2, dl_col3 = st.columns(3)

            with dl_col1:
                if export_format.startswith("Excel"):
                    data_bytes = to_excel_bytes(sheets)
                    fname = f"{base_name}_with_coordinates.xlsx"
                    mime = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                elif export_format.startswith("CSV"):
                    data_bytes = to_csv_bytes(sheets)
                    fname = f"{base_name}_with_coordinates.csv"
                    mime = "text/csv"
                else:
                    data_bytes = to_kml_bytes(sheets)
                    fname = f"{base_name}_with_coordinates.kml"
                    mime = "application/vnd.google-earth.kml+xml"

                st.download_button(
                    "⬇️ ดาวน์โหลดไฟล์พร้อมพิกัด",
                    data=data_bytes,
                    file_name=fname,
                    mime=mime,
                    use_container_width=True,
                )

            unmatched_df = pd.DataFrame(unmatched).drop_duplicates(subset=['Cust Code']) if unmatched else pd.DataFrame()

            with dl_col2:
                if not unmatched_df.empty:
                    unmatched_bytes = unmatched_to_excel_bytes(unmatched_df)
                    st.download_button(
                        f"📋 ดาวน์โหลดรายการที่ยังไม่มีพิกัด ({len(unmatched_df)})",
                        data=unmatched_bytes,
                        file_name=f"unmatched_{datetime.now().strftime('%Y%m%d')}.xlsx",
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        use_container_width=True,
                    )
                else:
                    st.markdown("""
                    <div class="rcm-card rcm-card-success" style="text-align:center;">
                        ✅ ไม่มีรายการตกหล่น — พิกัดครบทุกแถว
                    </div>
                    """, unsafe_allow_html=True)

            with dl_col3:
                total_unique_stops = sum(r.get('unique_stops', 0) for r in report)
                if unique_stops_sheets and total_unique_stops > 0:
                    stops_bytes = to_excel_bytes(unique_stops_sheets)
                    st.download_button(
                        f"📍 จุดส่งไม่ซ้ำ สำหรับ OptimoRoute ({total_unique_stops})",
                        data=stops_bytes,
                        file_name=f"{base_name}_unique_stops.xlsx",
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        use_container_width=True,
                        help="คอลัมน์เหมือนไฟล์หลักทุกอย่าง เพิ่มแค่คอลัมน์ 'จำนวนออเดอร์รวม' — ยุบแถวที่พิกัดซ้ำกันให้เหลือ 1 แถวต่อ 1 จุด",
                    )

            if not unmatched_df.empty:
                st.write("")
                st.markdown(f"""
                <div class="rcm-card rcm-card-warn">
                    ⚠️ <b>มี {len(unmatched_df)} รายการที่ยังไม่มีพิกัด</b> — ใช้ไฟล์ที่ดาวน์โหลดด้านบนไปเพิ่มลงใน Master Data ได้เลย (ที่ Google Sheets โดยตรง)
                </div>
                """, unsafe_allow_html=True)
                st.dataframe(unmatched_df, use_container_width=True, hide_index=True)

            ambiguous_df = pd.DataFrame(ambiguous).drop_duplicates(subset=['Cust Code']) if ambiguous else pd.DataFrame()
            if not ambiguous_df.empty:
                st.write("")
                st.markdown(f"""
                <div class="rcm-card rcm-card-warn">
                    🔀 <b>พบ {len(ambiguous_df)} รายการที่ชื่อซ้ำในหลายสาขา (พิกัดขัดแย้งกัน)</b> —
                    ระบบไม่กล้าเดาว่าเป็นสาขาไหน กรุณาเปิด Master Data แล้วแก้ Ship To Name ให้ระบุสาขาชัดเจน (เช่น เพิ่มชื่อสาขาต่อท้าย) แล้วลองประมวลผลใหม่
                </div>
                """, unsafe_allow_html=True)
                st.dataframe(ambiguous_df, use_container_width=True, hide_index=True)
                ambiguous_bytes = unmatched_to_excel_bytes(ambiguous_df)
                st.download_button(
                    f"🔀 ดาวน์โหลดรายการชื่อซ้ำที่ต้องตรวจสอบ ({len(ambiguous_df)})",
                    data=ambiguous_bytes,
                    file_name=f"ambiguous_{datetime.now().strftime('%Y%m%d')}.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    use_container_width=True,
                )

    st.markdown('<div class="rcm-glow-divider"></div>', unsafe_allow_html=True)
    st.caption(f"🛰️ Master Data sync (cache): {datetime.now().strftime('%Y-%m-%d %H:%M')} — รีเฟรชอัตโนมัติทุก 5 นาที หรือกด \"ซิงค์ตอนนี้\" ด้านบนเพื่อดึงข้อมูลล่าสุดทันที")

def _render_weight():
    if not PRODUCT_SHEET_CSV_URL:
        st.markdown("""
        <div class="rcm-card rcm-card-warn">
            ⚠️ <b>ยังไม่ได้ตั้งค่า PRODUCT_SHEET_CSV_URL</b><br>
            เพิ่มใน Streamlit Secrets: <code>PRODUCT_SHEET_CSV_URL = "https://docs.google.com/spreadsheets/d/YOUR_ID/export?format=csv&gid=0"</code>
        </div>
        """, unsafe_allow_html=True)
    else:
        try:
            product_df = load_product_master(PRODUCT_SHEET_CSV_URL)
            pstatus_col, psync_col = st.columns([5, 1])
            with pstatus_col:
                st.markdown(f"""
                <div class="rcm-status-card">
                    <div class="rcm-status-left">
                        <div class="rcm-status-icon">✅</div>
                        <div>
                            <div class="rcm-status-title">เชื่อมต่อ Product Master สำเร็จ</div>
                            <div class="rcm-status-sub">{len(product_df):,} รายการสินค้าในระบบ</div>
                        </div>
                    </div>
                    <div class="rcm-status-right">
                        <div class="rcm-status-right-icon">📦</div>
                        <div class="rcm-status-pill"><span class="rcm-status-dot"></span>Connected</div>
                    </div>
                </div>
                """, unsafe_allow_html=True)
            with psync_col:
                st.markdown('<div class="rcm-sync-btn-wrap">', unsafe_allow_html=True)
                if st.button("🔄 ซิงค์ตอนนี้", use_container_width=True, key="product_sync_btn",
                             help="ดึงข้อมูล Product Master ล่าสุดจาก Google Sheets ทันที"):
                    load_product_master.clear()
                    st.rerun()
                st.markdown('</div>', unsafe_allow_html=True)
        except Exception as e:
            st.error(f"โหลด Product Master ไม่สำเร็จ: {e}")
            return

        st.write("")
        search_term = st.text_input(
            "ค้นหาสินค้า", label_visibility="collapsed",
            placeholder="🔍 ค้นหาด้วย SKUCode หรือชื่อสินค้า...",
        )

        display_cols = ['SKUCode', 'PRODUCT', 'BRAND', 'CATEGORY', 'UOM',
                         'Pieces per Cases/Carton', 'Unit Net Wt (kg)',
                         'Case Weight (kg)', 'Case Volume (CBM)', 'Status']
        display_cols = [c for c in display_cols if c in product_df.columns]
        display_df = product_df[display_cols]

        if search_term:
            mask = (
                display_df['SKUCode'].astype(str).str.contains(search_term, case=False, na=False) |
                display_df['PRODUCT'].astype(str).str.contains(search_term, case=False, na=False)
            )
            display_df = display_df[mask]

        st.dataframe(display_df, use_container_width=True, hide_index=True, height=460)
        st.caption(f"แสดง {len(display_df):,} จาก {len(product_df):,} รายการทั้งหมด")

        st.markdown('<div class="rcm-glow-divider"></div>', unsafe_allow_html=True)
        st.markdown("##### ⚖️ คำนวณน้ำหนักจากไฟล์ route")
        st.caption("โยนไฟล์ route เข้ามา ระบบจะเชื่อม Product Code กับ Product Master แล้วเติมคอลัมน์น้ำหนักท้ายทุกแถว (คอลัมน์เดิมครบเหมือนไฟล์ต้นฉบับ)")

        weight_file = st.file_uploader(
            "อัพโหลดไฟล์ route สำหรับคำนวณน้ำหนัก", type=['xls', 'xlsx'],
            label_visibility="collapsed", key="weight_file_uploader",
        )

        if weight_file is not None:
            wfile_bytes = weight_file.read()
            try:
                wxl_peek = pd.ExcelFile(io.BytesIO(wfile_bytes))
                w_sheet_names = wxl_peek.sheet_names
            except Exception as e:
                st.error(f"เปิดไฟล์ไม่ได้: {e}")
                return

            st.write("")
            wsheet_cols = st.columns(min(len(w_sheet_names), 4))
            w_selected_sheets = []
            for i, sn in enumerate(w_sheet_names):
                with wsheet_cols[i % len(wsheet_cols)]:
                    checked = st.checkbox(sn, value=True, key=f"wsheet_{sn}")
                    if checked:
                        w_selected_sheets.append(sn)

            if not w_selected_sheets:
                st.warning("⚠️ กรุณาเลือกอย่างน้อย 1 Sheet เพื่อเริ่มประมวลผล")
                return

            st.write("")
            weight_process_clicked = st.button("⚖️ คำนวณน้ำหนัก", type="primary", use_container_width=True, key="weight_process_btn")

            if weight_process_clicked:
                with st.spinner("⚖️ กำลังคำนวณน้ำหนัก..."):
                    try:
                        w_sheets, w_report, w_no_weight = process_weight_file(wfile_bytes, product_df, w_selected_sheets)
                        st.session_state['last_w_sheets'] = w_sheets
                        st.session_state['last_w_report'] = w_report
                        st.session_state['last_w_no_weight'] = w_no_weight
                        st.session_state['last_w_filename'] = weight_file.name
                    except Exception as e:
                        st.error(f"เกิดข้อผิดพลาด: {e}")
                        st.exception(e)

            if 'last_w_sheets' in st.session_state and st.session_state.get('last_w_filename') == weight_file.name:
                w_sheets = st.session_state['last_w_sheets']
                w_report = st.session_state['last_w_report']
                w_no_weight = st.session_state['last_w_no_weight']

                w_total_calc = sum(r['calculated'] for r in w_report)
                w_total_rows = sum(r['total'] for r in w_report)
                w_rate = (w_total_calc / w_total_rows * 100) if w_total_rows else 0

                w_card_class = "rcm-card-success" if w_rate >= 90 else "rcm-card-warn"
                st.markdown(f"""
                <div class="rcm-card {w_card_class}">
                    ⚖️ <b>เสร็จแล้ว!</b> คำนวณน้ำหนักได้ {w_total_calc:,} / {w_total_rows:,} แถว ({w_rate:.1f}%)
                </div>
                """, unsafe_allow_html=True)

                w_cols = st.columns(len(w_report))
                for col, r in zip(w_cols, w_report):
                    with col:
                        if r['skipped']:
                            stat_box(f"SHEET: {r['sheet']}", "ข้าม", "ไม่พบคอลัมน์ Product Code")
                        else:
                            stat_box(f"SHEET: {r['sheet']}", f"{r['calculated']}/{r['total']}", "คำนวณสำเร็จ")

                st.write("")
                w_dl_col1, w_dl_col2 = st.columns(2)
                w_base_name = weight_file.name.rsplit('.', 1)[0]

                with w_dl_col1:
                    w_output_bytes = to_excel_bytes(w_sheets)
                    st.download_button(
                        "⬇️ ดาวน์โหลดไฟล์พร้อมน้ำหนัก",
                        data=w_output_bytes,
                        file_name=f"{w_base_name}_with_weight.xlsx",
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        use_container_width=True,
                        key="weight_download_main",
                    )

                w_no_weight_df = pd.DataFrame(w_no_weight).drop_duplicates(subset=['Product Code', 'สาเหตุ']) if w_no_weight else pd.DataFrame()

                with w_dl_col2:
                    if not w_no_weight_df.empty:
                        w_nw_bytes = unmatched_to_excel_bytes(w_no_weight_df)
                        st.download_button(
                            f"📋 รายการที่ยังไม่มีน้ำหนัก ({len(w_no_weight_df)})",
                            data=w_nw_bytes,
                            file_name=f"no_weight_{datetime.now().strftime('%Y%m%d')}.xlsx",
                            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                            use_container_width=True,
                            key="weight_download_missing",
                        )
                    else:
                        st.markdown("""
                        <div class="rcm-card rcm-card-success" style="text-align:center;">
                            ✅ คำนวณน้ำหนักได้ครบทุกแถว
                        </div>
                        """, unsafe_allow_html=True)

                if not w_no_weight_df.empty:
                    st.write("")
                    st.markdown(f"""
                    <div class="rcm-card rcm-card-warn">
                        ⚠️ <b>มี {len(w_no_weight_df)} รายการที่ยังคำนวณน้ำหนักไม่ได้</b> — ดูคอลัมน์ "สาเหตุ" เพื่อรู้ว่าต้องแก้อะไร (ส่วนใหญ่คือ Product Master ยังไม่มีข้อมูลน้ำหนักของ SKU นั้น)
                    </div>
                    """, unsafe_allow_html=True)
                    st.dataframe(w_no_weight_df, use_container_width=True, hide_index=True)


# ============================================================
# TAB 3: QR CODE GENERATOR
# ============================================================
def generate_qr_bytes(link, box_size=10, border=4, fill_color='#000000', back_color='#FFFFFF'):
    qr = qrcode.QRCode(
        version=None,
        error_correction=qrcode.constants.ERROR_CORRECT_M,
        box_size=box_size,
        border=border,
    )
    qr.add_data(link)
    qr.make(fit=True)
    img = qr.make_image(fill_color=fill_color, back_color=back_color).convert('RGB')
    buf = io.BytesIO()
    img.save(buf, format='PNG')
    return buf.getvalue(), img.size


# ---- Batch mode: read links straight out of an uploaded file, embed a QR per row ----
QR_URL_RE = re.compile(r'https?://\S+')


def find_link_column(sheet_rows):
    """sheet_rows: list of row tuples (raw cell values, 0-indexed).
    Scans every cell (content, not header text) and returns the 0-based column
    index with the most URL-looking values — or None if the sheet has none.
    Content-based on purpose: header names for a maps-link column vary a lot
    between files ('Ship to Google Plus code (Epicor)', 'ที่ตั้ง', a column with
    no header at all, ...), so we trust what's actually in the cells."""
    counts = {}
    for row in sheet_rows:
        for i, v in enumerate(row):
            if isinstance(v, str) and QR_URL_RE.search(v):
                counts[i] = counts.get(i, 0) + 1
    if not counts:
        return None
    return max(counts, key=counts.get)


def _guess_table_top(raw_rows):
    """หาแถวหัวตาราง (0-indexed): แถวแรกที่มีข้อมูลเกินครึ่งของแถวที่กว้างที่สุด (อย่างน้อย 3 ช่อง)
    ใช้แค่กำหนดตำแหน่งวาง QR — ถ้าหาไม่เจอคืน None"""
    counts = [sum(1 for v in row if not (v is None or (isinstance(v, float) and pd.isna(v)) or str(v).strip() == ''))
              for row in raw_rows]
    if not counts:
        return None
    need = max(3, max(counts) / 2)
    for i, c in enumerate(counts):
        if c >= need:
            return i
    return None


def _embed_qr_column(ws, raw_rows, link_col, start_col0, box_size=5, border=2):
    """Appends a 'QR Code' column at 0-based column index start_col0, embedding
    a QR image at NATIVE resolution (no shrinking) for every row whose link_col
    cell contains a link — long multi-stop Google Maps links need a fairly
    dense/large QR to stay scannable, so images are never downscaled to a fixed
    thumbnail. Never touches any other existing cell. Returns qr_count."""
    new_col_letter = get_column_letter(start_col0 + 1)
    ws.cell(row=1, column=start_col0 + 1, value="QR Code")

    links = []
    for r_idx, row in enumerate(raw_rows):
        v = row[link_col] if link_col < len(row) else None
        if isinstance(v, str):
            m = QR_URL_RE.search(v)
            if m:
                links.append((r_idx, m.group(0)))

    # ลิงก์เดียวทั้งชีท (เช่นลิงก์เส้นทางท้ายชีทของไฟล์ OptimoRoute) -> วาง QR ลอยไว้ขวาบนข้างตาราง
    # เริ่มที่แถวข้อมูลแรก และไม่ยืดความสูงแถวไหนเลย (แถวลิงก์ด้านล่างสูงปกติ)
    # ลิงก์หลายแถว (เช่น Master Data) -> วาง QR ข้างแถวของมันเหมือนเดิม จะได้รู้ว่า QR ไหนของแถวไหน
    if len(links) == 1:
        r_idx, url = links[0]
        try:
            qr_bytes, (w_px, h_px) = generate_qr_bytes(url, box_size=box_size, border=border)
        except Exception:
            return 0
        top = _guess_table_top(raw_rows)
        anchor_row = min(top + 2, r_idx + 1) if top is not None else 2  # แถวข้อมูลแรก (1-indexed)
        ws.add_image(XLImage(io.BytesIO(qr_bytes)), f"{new_col_letter}{anchor_row}")
        ws.column_dimensions[new_col_letter].width = max(12, w_px / 7 + 2)
        return 1

    qr_count = 0
    max_w_px = 0
    for r_idx, url in links:
        try:
            qr_bytes, (w_px, h_px) = generate_qr_bytes(url, box_size=box_size, border=border)
        except Exception:
            continue
        img = XLImage(io.BytesIO(qr_bytes))
        excel_row = r_idx + 1  # 1-indexed
        ws.add_image(img, f"{new_col_letter}{excel_row}")
        target_h = h_px * 0.75 + 4  # px -> pt, + a little padding
        current_h = ws.row_dimensions[excel_row].height
        if not current_h or current_h < target_h:
            ws.row_dimensions[excel_row].height = target_h
        max_w_px = max(max_w_px, w_px)
        qr_count += 1

    if qr_count:
        ws.column_dimensions[new_col_letter].width = max(12, max_w_px / 7 + 2)

    return qr_count


# ---- "Keep template": preserve original .xls formatting (font, fill, border,
# column widths, row heights, merged cells) when rebuilding a legacy .xls into
# .xlsx, instead of a bare values-only rebuild. Only needed for the .xls
# fallback — real .xlsx already keeps 100% of its original formatting since
# that path edits the loaded workbook in place. ----
_QR_BORDER_STYLE_MAP = {
    0: None, 1: 'thin', 2: 'medium', 3: 'dashed', 4: 'dotted',
    5: 'thick', 6: 'double', 7: 'hair', 8: 'mediumDashed',
    9: 'dashDot', 10: 'mediumDashDot', 11: 'dashDotDot',
    12: 'mediumDashDotDot', 13: 'slantDashDot',
}
_QR_HALIGN_MAP = {1: 'left', 2: 'center', 3: 'right', 4: 'fill', 5: 'justify', 6: 'centerContinuous'}
_QR_VALIGN_MAP = {0: 'top', 1: 'center', 2: 'bottom', 3: 'justify'}


def _xlrd_color_hex(color_map, idx):
    """xlrd colour index -> 'RRGGBB' hex, or None for an 'automatic/default'
    color (index not in the palette — leave openpyxl's own default alone)."""
    if idx is None:
        return None
    rgb = color_map.get(idx)
    return ('%02X%02X%02X' % rgb) if rgb else None


def _xlrd_border_side(style_code, colour_idx, color_map):
    if not style_code:
        return Side(style=None)
    style = _QR_BORDER_STYLE_MAP.get(style_code, 'thin')
    color_hex = _xlrd_color_hex(color_map, colour_idx)
    return Side(style=style, color=color_hex) if color_hex else Side(style=style)


def _build_openpyxl_style_from_xf(xf, font_rec, color_map):
    """One xlrd XF (+ its font) record -> (Font, PatternFill, Border, Alignment)."""
    font = Font(
        name=font_rec.name or 'Calibri',
        size=(font_rec.height / 20.0) if font_rec.height else 11,
        bold=bool(font_rec.bold),
        italic=bool(font_rec.italic),
        color=_xlrd_color_hex(color_map, font_rec.colour_index),
    )
    fill = PatternFill()
    if getattr(xf.background, 'fill_pattern', 0):
        bg_hex = _xlrd_color_hex(color_map, xf.background.pattern_colour_index)
        if bg_hex:
            fill = PatternFill(fill_type='solid', fgColor=bg_hex)
    border = Border(
        top=_xlrd_border_side(xf.border.top_line_style, xf.border.top_colour_index, color_map),
        bottom=_xlrd_border_side(xf.border.bottom_line_style, xf.border.bottom_colour_index, color_map),
        left=_xlrd_border_side(xf.border.left_line_style, xf.border.left_colour_index, color_map),
        right=_xlrd_border_side(xf.border.right_line_style, xf.border.right_colour_index, color_map),
    )
    alignment = Alignment(
        horizontal=_QR_HALIGN_MAP.get(xf.alignment.hor_align),
        vertical=_QR_VALIGN_MAP.get(xf.alignment.vert_align),
        wrap_text=bool(getattr(xf.alignment, 'wrap', False)),
    )
    return font, fill, border, alignment


def _apply_xls_styles(ws, sh, xlrd_book, n_rows, n_cols):
    """Layers font/fill/border/alignment + column widths + row heights +
    merged cells from an xlrd(formatting_info=True) sheet onto an already
    value-populated openpyxl worksheet ws — bounded to [0,n_rows)x[0,n_cols),
    the same real-data bounds pandas already determined. That bound matters:
    some real CCS .xls exports report tens of thousands of blank padding rows
    at the raw-file level (e.g. a stray format applied far past the real data)
    — capping the loop at pandas' trimmed shape keeps this fast and correct
    regardless of what the raw file claims."""
    color_map = xlrd_book.colour_map
    xf_list = xlrd_book.xf_list
    font_list = xlrd_book.font_list
    style_cache = {}

    r_bound = min(n_rows, sh.nrows)
    c_bound = min(n_cols, sh.ncols)

    for r in range(r_bound):
        for c in range(c_bound):
            try:
                xf_index = sh.cell_xf_index(r, c)
            except IndexError:
                continue
            if xf_index not in style_cache:
                try:
                    xf = xf_list[xf_index]
                    font_rec = font_list[xf.font_index]
                    style_cache[xf_index] = _build_openpyxl_style_from_xf(xf, font_rec, color_map)
                except Exception:
                    style_cache[xf_index] = None
            style = style_cache[xf_index]
            if style is None:
                continue
            font, fill, border, alignment = style
            wc = ws.cell(row=r + 1, column=c + 1)
            wc.font = font
            wc.fill = fill
            wc.border = border
            wc.alignment = alignment

    for c in range(c_bound):
        ci = sh.colinfo_map.get(c)
        if ci and ci.width:
            ws.column_dimensions[get_column_letter(c + 1)].width = ci.width / 256.0

    for r in range(r_bound):
        ri = sh.rowinfo_map.get(r)
        if ri and ri.height:
            ws.row_dimensions[r + 1].height = ri.height / 20.0

    for (rlo, rhi, clo, chi) in sh.merged_cells:
        if rlo < r_bound and clo < c_bound:
            try:
                ws.merge_cells(start_row=rlo + 1, end_row=min(rhi, r_bound),
                                start_column=clo + 1, end_column=min(chi, c_bound))
            except Exception:
                pass


_HYPERLINK_FORMULA_RE = re.compile(
    r'^=\s*HYPERLINK\(\s*"((?:[^"]|"")*)"\s*(?:[,;]\s*"((?:[^"]|"")*)"\s*)?\)\s*$',
    re.IGNORECASE,
)


def _flatten_hyperlink_formulas(ws, cached_ws=None):
    """=HYPERLINK("url","ชื่อ") -> ค่าเซลล์ = "ชื่อ" (ข้อความธรรมดา) + cell.hyperlink = url
    ชื่อที่ใช้: ค่าที่ Excel คำนวณเก็บไว้ (ถ้ามี) ไม่งั้นอ่านจากอาร์กิวเมนต์ตัวที่ 2 ในสูตร
    ถ้าสูตรซับซ้อนกว่านี้ (อ้างเซลล์อื่น ฯลฯ) และไม่มีค่าที่คำนวณไว้ -> ปล่อยสูตรเดิมไว้ ไม่เดา
    Returns จำนวนเซลล์ที่แปลง"""
    n = 0
    for row in ws.iter_rows():
        for cell in row:
            v = cell.value
            if not (isinstance(v, str) and v.lstrip().upper().startswith('=HYPERLINK(')):
                continue
            m = _HYPERLINK_FORMULA_RE.match(v.strip())
            url = m.group(1).replace('""', '"') if m else None
            name = m.group(2).replace('""', '"') if (m and m.group(2) is not None) else None
            cached = cached_ws[cell.coordinate].value if cached_ws is not None else None
            display = cached if cached not in (None, '') else (name if name is not None else url)
            if display is None:
                continue
            cell.value = display
            if url:
                cell.hyperlink = url
            n += 1
    return n


def process_qr_batch_file(file_bytes, selected_sheets=None, box_size=5, border=2):
    """
    Scans each selected sheet for a column containing links and returns a NEW
    workbook (bytes) with a 'QR Code' column appended to every sheet that has
    one — one embedded, natively-scannable QR image per row with a link.

    Never alters or removes any existing cell — pure append. A sheet with no
    detected link is left completely untouched (not even a header added).

    Returns (output_bytes, report): report is a list of dicts per sheet —
    {'sheet', 'link_col_letter' (or None), 'qr_count', 'total_rows'}.
    """
    xl = pd.ExcelFile(io.BytesIO(file_bytes))
    all_sheet_names = xl.sheet_names
    sheets_to_process = selected_sheets if selected_sheets else all_sheet_names

    # High-fidelity path: real .xlsx loads straight into openpyxl, so every
    # other sheet, style, formula and column is preserved byte-for-byte and we
    # only ever append the new column. Legacy .xls can't be loaded (or written)
    # by openpyxl, so it falls back to a values-only rebuild — same as the rest
    # of this app already does for .xls uploads.
    try:
        out_wb = load_workbook(io.BytesIO(file_bytes))
        native = True
    except Exception:
        out_wb = None
        native = False

    report = []

    if native:
        # เซลล์ที่เป็นสูตร =HYPERLINK(...) (เช่นคอลัมน์ "ที่ตั้ง" ในไฟล์ OptimoRoute) -> เขียนชื่อจริงลงไป
        # เป็นข้อความธรรมดา + ใส่ลิงก์จริงไว้ที่เซลล์ (ยังกดได้) ทำก่อนสแกนหาลิงก์ เพื่อไม่ให้ข้อความสูตร
        # ที่มี https:// ถูกเข้าใจผิดว่าเป็นคอลัมน์ลิงก์
        try:
            cached_wb = load_workbook(io.BytesIO(file_bytes), data_only=True)
        except Exception:
            cached_wb = None
        for ws_any in out_wb.worksheets:
            cached_ws = cached_wb[ws_any.title] if cached_wb is not None and ws_any.title in cached_wb.sheetnames else None
            _flatten_hyperlink_formulas(ws_any, cached_ws)

        for sn in sheets_to_process:
            if sn not in out_wb.sheetnames:
                continue
            ws = out_wb[sn]
            raw_rows = list(ws.iter_rows(values_only=True))
            link_col = find_link_column(raw_rows)

            if link_col is None:
                report.append({'sheet': sn, 'link_col_letter': None, 'qr_count': 0, 'total_rows': len(raw_rows)})
                continue

            qr_count = _embed_qr_column(ws, raw_rows, link_col, ws.max_column, box_size, border)
            report.append({
                'sheet': sn,
                'link_col_letter': get_column_letter(link_col + 1),
                'qr_count': qr_count,
                'total_rows': len(raw_rows),
            })

        buf = io.BytesIO()
        out_wb.save(buf)
        return buf.getvalue(), report

    else:
        # Legacy .xls: values still come from the same pandas path as before
        # (unchanged — already trims fake padding ranges and handles date/time
        # cells correctly). "Keep template": layer the original font/fill/
        # border/column-width/merge formatting back on top via a separate
        # xlrd(formatting_info=True) pass. If that pass fails for any reason,
        # degrade silently to the old values-only look rather than breaking
        # the actual QR feature over a styling problem.
        try:
            xlrd_book = xlrd.open_workbook(file_contents=file_bytes, formatting_info=True)
        except Exception:
            xlrd_book = None

        out_wb = Workbook()
        out_wb.remove(out_wb.active)

        for sn in sheets_to_process:
            if sn not in all_sheet_names:
                continue
            df_raw = pd.read_excel(xl, sheet_name=sn, header=None)
            raw_rows = df_raw.values.tolist()

            ws = out_wb.create_sheet(title=sn[:31])
            for r_idx, row in enumerate(raw_rows):
                for c_idx, v in enumerate(row):
                    if pd.isna(v):
                        continue
                    ws.cell(row=r_idx + 1, column=c_idx + 1, value=v)

            if xlrd_book is not None:
                try:
                    _apply_xls_styles(ws, xlrd_book.sheet_by_name(sn), xlrd_book, df_raw.shape[0], df_raw.shape[1])
                except Exception:
                    pass

            link_col = find_link_column(raw_rows)
            if link_col is None:
                report.append({'sheet': sn, 'link_col_letter': None, 'qr_count': 0, 'total_rows': len(raw_rows)})
                continue

            qr_count = _embed_qr_column(ws, raw_rows, link_col, df_raw.shape[1], box_size, border)
            report.append({
                'sheet': sn,
                'link_col_letter': get_column_letter(link_col + 1),
                'qr_count': qr_count,
                'total_rows': len(raw_rows),
            })

        buf = io.BytesIO()
        out_wb.save(buf)
        return buf.getvalue(), report


def _render_qr():
    st.markdown("##### 🔗 สร้าง QR Code จากลิงก์")
    st.caption("ใส่ลิงก์อะไรก็ได้ — ลิงก์แอป, Google Sheets, เอกสาร ฯลฯ — ได้ QR Code กลับมาดาวน์โหลดได้ทันที")

    qr_link = st.text_input(
        "ลิงก์", label_visibility="collapsed",
        placeholder="https://example.com",
        key="qr_link_input",
    )

    with st.expander("⚙️ ปรับแต่ง (ไม่บังคับ)"):
        qc1, qc2 = st.columns(2)
        with qc1:
            qr_size = st.slider("ขนาด QR Code", min_value=5, max_value=20, value=10, key="qr_size")
            qr_fill = st.color_picker("สีลาย", value="#000000", key="qr_fill")
        with qc2:
            qr_border = st.slider("ขอบขาว (border)", min_value=1, max_value=10, value=4, key="qr_border")
            qr_back = st.color_picker("สีพื้นหลัง", value="#FFFFFF", key="qr_back")

    if qr_link and qr_link.strip():
        try:
            qr_bytes, qr_dims = generate_qr_bytes(
                qr_link.strip(),
                box_size=qr_size, border=qr_border,
                fill_color=qr_fill, back_color=qr_back,
            )
            st.write("")
            qr_img_col, qr_info_col = st.columns([1, 1])
            with qr_img_col:
                st.image(qr_bytes, width=280)
            with qr_info_col:
                st.markdown(f"""
                <div class="rcm-card rcm-card-success">
                    ✅ <b>สร้าง QR Code สำเร็จ</b>
                    <div class="rcm-status-sub" style="margin-top:6px;">ขนาดภาพ: {qr_dims[0]}×{qr_dims[1]} px</div>
                </div>
                """, unsafe_allow_html=True)
                st.download_button(
                    "⬇️ ดาวน์โหลด QR Code (PNG)",
                    data=qr_bytes,
                    file_name="qrcode.png",
                    mime="image/png",
                    use_container_width=True,
                    key="qr_download_btn",
                )
        except Exception as e:
            st.error(f"สร้าง QR Code ไม่สำเร็จ: {e}")
    else:
        st.markdown("""
        <div class="rcm-card" style="text-align:center; color:#6B84A6;">
            📎 ใส่ลิงก์ด้านบนเพื่อเริ่มสร้าง QR Code
        </div>
        """, unsafe_allow_html=True)

    st.markdown('<div class="rcm-glow-divider"></div>', unsafe_allow_html=True)

    # ============================================================
    # BATCH: อ่านลิงก์จากไฟล์ที่อัพโหลด → ได้ไฟล์เดิมที่ฝัง QR Code กลับมา
    # ============================================================
    st.markdown("##### 📁 สร้าง QR Code จากไฟล์ (Batch)")
    st.caption("อัพโหลดไฟล์ที่มีลิงก์ Google Maps อยู่แล้ว (เช่นไฟล์ route ที่มีลิงก์เส้นทางท้ายชีท หรือ Master Data ที่มีลิงก์ต่อแถว) ระบบจะหาคอลัมน์ที่มีลิงก์ให้เอง แล้วฝัง QR Code กลับเข้าไปในคอลัมน์ใหม่ท้ายไฟล์ — ข้อมูลเดิมไม่ถูกแก้หรือลบแม้แต่แถวเดียว")

    qr_file = st.file_uploader("อัพโหลดไฟล์", type=['xls', 'xlsx'], label_visibility="collapsed", key="qr_batch_uploader")

    if qr_file is not None:
        qr_file_bytes = qr_file.read()

        try:
            qr_xl_peek = pd.ExcelFile(io.BytesIO(qr_file_bytes))
            qr_all_sheet_names = qr_xl_peek.sheet_names
        except Exception as e:
            st.error(f"เปิดไฟล์ไม่ได้: {e}")
            return

        st.markdown("###### 📑 เลือก Sheet ที่ต้องการสแกนหาลิงก์")
        qr_sheet_cols = st.columns(min(len(qr_all_sheet_names), 4))
        qr_selected_sheets = []
        for i, sn in enumerate(qr_all_sheet_names):
            with qr_sheet_cols[i % len(qr_sheet_cols)]:
                checked = st.checkbox(sn, value=True, key=f"qrbatch_sheet_{sn}")
                if checked:
                    qr_selected_sheets.append(sn)

        if not qr_selected_sheets:
            st.warning("⚠️ กรุณาเลือกอย่างน้อย 1 Sheet")
            return

        st.write("")
        qr_batch_clicked = st.button("🔍 สแกนหาลิงก์ + สร้าง QR Code", type="primary", use_container_width=True, key="qr_batch_btn")

        if qr_batch_clicked:
            with st.spinner("🔍 กำลังสแกนหาลิงก์และสร้าง QR Code..."):
                try:
                    qr_out_bytes, qr_report = process_qr_batch_file(qr_file_bytes, qr_selected_sheets)
                    st.session_state['last_qr_batch_bytes'] = qr_out_bytes
                    st.session_state['last_qr_batch_report'] = qr_report
                    st.session_state['last_qr_batch_filename'] = qr_file.name
                except Exception as e:
                    st.error(f"เกิดข้อผิดพลาด: {e}")
                    st.exception(e)

        if 'last_qr_batch_report' in st.session_state and st.session_state.get('last_qr_batch_filename') == qr_file.name:
            qr_report = st.session_state['last_qr_batch_report']
            qr_out_bytes = st.session_state['last_qr_batch_bytes']
            total_qr = sum(r['qr_count'] for r in qr_report)

            if total_qr > 0:
                st.markdown(f"""
                <div class="rcm-card rcm-card-success">
                    🎉 <b>เสร็จแล้ว!</b> สร้าง QR Code ได้ {total_qr:,} อัน
                </div>
                """, unsafe_allow_html=True)
            else:
                st.markdown("""
                <div class="rcm-card rcm-card-warn">
                    ⚠️ ไม่พบลิงก์ในไฟล์นี้เลย — ลองเช็คว่ามีคอลัมน์ที่เก็บลิงก์ Google Maps อยู่จริงไหม
                </div>
                """, unsafe_allow_html=True)

            qr_report_cols = st.columns(len(qr_report))
            for col, r in zip(qr_report_cols, qr_report):
                with col:
                    if r['link_col_letter'] is None:
                        stat_box(f"SHEET: {r['sheet']}", "ไม่พบลิงก์", f"{r['total_rows']:,} แถว")
                    else:
                        stat_box(f"SHEET: {r['sheet']}", f"{r['qr_count']:,} QR", f"พบลิงก์ในคอลัมน์ {r['link_col_letter']}")

            if total_qr > 0:
                st.write("")
                qr_batch_base_name = qr_file.name.rsplit('.', 1)[0]
                st.download_button(
                    "⬇️ ดาวน์โหลดไฟล์ที่ฝัง QR Code แล้ว (Excel)",
                    data=qr_out_bytes,
                    file_name=f"{qr_batch_base_name}_with_qr.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    use_container_width=True,
                    key="qr_batch_download_btn",
                )


# ============================================================
# TAB 4: FOR BKK — Sheet1 -> ใบงานต่อออเดอร์ (แยก Retail และ InterExp/B&W ออกเป็นชีทของตัวเอง)
# ============================================================
BKK_SERVICE_COL = 'Service Time (นาที)'
# ท้ายใบงาน (ลอกข้อความจากใบงาน Sheet2 ของทีม)
BKK_FOOTER_SIGN = ('ลงชื่อ_______________________________พนักงานขับ', 'รถเวลา________________________ออก')
BKK_FOOTER_NOTES = [
    'ถึงพนักงานขับรถ ถ้าในรูทมี สินค้า 4AOB ทุกตัว ให้เอากล่องสินค้าแพ็คลงกล่องโฟม และใส่ดรายไอซ์กล่องละ 2 ก้อน ก่อนนำสินค้าไปส่งลูกค้า เพราะขนมเสียหายง่ายต้อง ***',
    'ถึงพนักงานขับรถ ถ้าในรูทมี สินค้า 4HFF ทุกตัว ให้เอากล่องสินค้าแพ็คลงกล่องโฟม และใส่ดรายไอซ์กล่องละ 2 ก้อน ก่อนนำสินค้าไปส่งลูกค้า เพราะสินค้าเสียหายง่ายต้อง ***',
]
# น้ำหนักต่อออเดอร์ (รวมจากทุกรายการสินค้าใน Sheet1 + Product Master) และจำนวนสินค้าที่ต้องใช้ Dry Ice
BKK_WEIGHT_COL = 'น้ำหนัก (kg)'
BKK_AOB_COL = '4AOB (กล่อง)'
BKK_HFF_COL = '4HFF (PCS)'
BKK_DRYICE_GROUPS = (('4AOB', BKK_AOB_COL), ('4HFF', BKK_HFF_COL))  # Product Code ขึ้นต้นด้วย -> คอลัมน์
BKK_COLUMNS = ['Order Date', 'OrderNo', 'Customers Code', 'Cust Type', 'Customer Name', 'ShipTo Name',
               'ShipTo.Address', 'Route', 'Remark', 'เวลาเข้า', 'เวลาออก', BKK_SERVICE_COL,
               BKK_WEIGHT_COL, BKK_AOB_COL, BKK_HFF_COL, 'Latitude', 'Longitude']
# ความกว้างคอลัมน์ลอกจากใบงาน Sheet2 ที่ทีมใช้อยู่จริง (A..K) + Service Time + น้ำหนัก/4AOB/4HFF + Lat/Lon
BKK_COL_WIDTHS = [9.9, 39.1, 14.4, 9.1, 53.3, 53.7, 40.1, 19.4, 27.4, 7.7, 8.6, 15.5, 12, 12, 11, 12.5, 12.5]


def build_channel_map(df):
    """Customer Master ทั้งตาราง (ไม่ตัดแถวที่ยังไม่มีพิกัด) -> {ShipTo Name ที่ normalize แล้ว: Channel}
    Channel คือคอลัมน์ C ของ Master; ถ้าชื่อเดียวกันมีหลายแถว ใช้แถวแรกที่เจอ"""
    ch_col = next((c for c in df.columns if str(c).strip().lower() == 'channel'), None)
    if ch_col is None:
        ch_col = df.columns[2]
    out = {}
    for name, ch in zip(df['Ship To Name'], df[ch_col]):
        key = normalize(name)
        if key and key not in out and not _cell_blank(ch):
            out[key] = str(ch).strip()
    return out


@st.cache_data(ttl=300)
def load_master_channels(url):
    df = pd.read_csv(url)
    df.columns = [c.strip() for c in df.columns]
    if 'Ship To Name' not in df.columns:
        raise ValueError("ไม่พบคอลัมน์ Ship To Name ใน Master")
    return build_channel_map(df)


def _bkk_clean_key(v):
    """975484.0 -> 975484, ' 5 ' -> '5', 1.0 -> 1 (ให้ OrderNo / Route เทียบกันได้ตรงๆ)"""
    if isinstance(v, float) and v == int(v):
        return int(v)
    if isinstance(v, str):
        s = v.strip()
        return int(s) if s.isdigit() else s
    return v


def read_bkk_orders(file_bytes):
    """อ่าน Sheet1 เท่านั้น -> 1 แถวต่อ 1 OrderNo (แถวแรกที่เจอของออเดอร์นั้น) ตามลำดับเดิมในไฟล์
    ตัดแถวว่างทิ้ง Returns (orders_df, info)"""
    xl = pd.ExcelFile(io.BytesIO(file_bytes))
    sheet = 'Sheet1' if 'Sheet1' in xl.sheet_names else xl.sheet_names[0]
    raw = xl.parse(sheet, header=None).values.tolist()
    header_row, c_order = find_header_row(raw, ['orderno', 'order no', 'order no.'])
    if header_row is None:
        raise ValueError(f'ไม่พบคอลัมน์ OrderNo ในชีท "{sheet}"')
    header = [v.strip().lower() if isinstance(v, str) else '' for v in raw[header_row]]

    def col(*names, prefix=None):
        for i, h in enumerate(header):
            if h in names:
                return i
        if prefix:
            for i, h in enumerate(header):
                if h.startswith(prefix):
                    return i
        return None

    c_date = col('order date')
    c_code = col('cust code', 'cust id', 'cust. id', 'cust.id', 'customers code', 'customer code')
    c_type = col('cust type', 'customer type')
    c_cname = col('cust name', 'cust. name', 'customer name')
    c_sname = col('shipto name', 'ship to name', 'ship-to name')
    c_addr = col('shipto.address', prefix='shipto.address')
    if c_addr is None:
        c_addr = col('shipto.c')
    c_route = col('route')
    c_prod = col('product code')
    c_qty = col('qty')
    c_unit = col('units', 'unit', 'uom')

    def get(row, c):
        if c is None or c >= len(row) or _cell_blank(row[c]):
            return None
        v = row[c]
        return v.strip() if isinstance(v, str) else v

    rows, seen, multi_route, lines = [], {}, {}, []
    line_count = no_order_rows = 0
    for row in raw[header_row + 1:]:
        if all(_cell_blank(v) for v in row):
            continue  # แถวว่าง
        if _cell_blank(row[c_order]):
            no_order_rows += 1  # มีข้อมูลแต่ไม่มีเลขออเดอร์ -> นับไว้รายงาน
            continue
        line_count += 1
        order_no = _bkk_clean_key(row[c_order])
        route = _bkk_clean_key(get(row, c_route))
        if get(row, c_prod) is not None:
            # เก็บรายการสินค้าไว้คิดน้ำหนัก / นับ 4AOB 4HFF ต่อออเดอร์
            lines.append({'OrderNo': order_no, 'Product Code': str(get(row, c_prod)).strip(),
                          'Qty': get(row, c_qty), 'Units': get(row, c_unit)})
        if order_no in seen:
            first = seen[order_no]
            if route != first['Route'] and route is not None and first['Route'] is not None:
                multi_route.setdefault(order_no, [first['Route']])
                if route not in multi_route[order_no]:
                    multi_route[order_no].append(route)
            # แถวแรกของออเดอร์นี้มีช่องไหนว่าง แต่แถวถัดไปของออเดอร์เดียวกันมีค่า -> เติมให้ (ไม่ทับค่าที่มีอยู่)
            for fld, c in (('Order Date', c_date), ('Customers Code', c_code), ('Cust Type', c_type),
                           ('Customer Name', c_cname), ('ShipTo Name', c_sname), ('ShipTo.Address', c_addr),
                           ('Route', c_route)):
                if first[fld] is None:
                    v = get(row, c)
                    if v is not None:
                        first[fld] = _bkk_clean_key(v) if fld == 'Route' else v
            continue
        rec = {
            'Order Date': get(row, c_date), 'OrderNo': order_no, 'Customers Code': get(row, c_code),
            'Cust Type': get(row, c_type), 'Customer Name': get(row, c_cname),
            'ShipTo Name': get(row, c_sname), 'ShipTo.Address': get(row, c_addr), 'Route': route,
        }
        seen[order_no] = rec
        rows.append(rec)

    orders = pd.DataFrame(rows, columns=BKK_COLUMNS[:8])
    info = {'sheet': sheet, 'lines': line_count, 'orders': len(orders), 'no_order_rows': no_order_rows,
            'multi_route': multi_route, 'has_shipto_col': c_sname is not None,
            'lines_df': pd.DataFrame(lines, columns=['OrderNo', 'Product Code', 'Qty', 'Units'])}
    return orders, info


def _bkk_route_sort_key(route):
    """Route ตัวเลขมาก่อนเรียงตามค่า (1, 2, 3, 11, 12) แล้วตามด้วย Route ที่เป็นตัวหนังสือ"""
    if isinstance(route, (int, float)) and not pd.isna(route):
        return (0, float(route), '')
    return (1, 0.0, '' if route is None else str(route))


BKK_CARRIER_SHEET = 'InterExp-B&W'
# ใบงาน: Retail กับ InterExp-B&W อยู่ชีทเดียวกัน (Retail ก่อน แล้วต่อด้วย InterExp-B&W) — Service Time ยังคิดแยกกลุ่มเหมือนเดิม
BKK_RETAIL_CARRIER_SHEET = 'Retail + InterExp-B&W'


def bkk_output_sheets(normal_df, retail_df, carrier_df, notype_df=None):
    """ชีทที่อยู่ในไฟล์ใบงาน: Order / Retail + InterExp-B&W / (No Cust Type ถ้ามี)"""
    sheets = [('Order', normal_df),
              (BKK_RETAIL_CARRIER_SHEET, pd.concat([retail_df, carrier_df], ignore_index=True))]
    if notype_df is not None and len(notype_df):
        sheets.append((BKK_NOTYPE_SHEET, notype_df))
    return sheets
BKK_NOTYPE_SHEET = 'No Cust Type'

# ลูกค้าที่ในระบบเป็น R (Retail) แต่ให้ถือเป็น F (Food Service) เสมอ — เพิ่มชื่อในลิสต์นี้ได้
BKK_FORCE_FOOD_SERVICE = ['FMR-for RT Sale Dept.']


def _bkk_name_key(v):
    """เทียบชื่อแบบไม่สนตัวพิมพ์ เว้นวรรค จุด ขีด: 'FMR- for RT Sale Dept' == 'FMR-for RT Sale Dept.'"""
    return re.sub(r'[^A-Z0-9ก-๙]', '', str(v).upper()) if v is not None else ''


def _bkk_is_carrier_route(route):
    """Route ที่ส่งผ่านขนส่งภายนอก: InterExp และ IE,B&W (เขียนได้หลายแบบ เช่น 'IE B&W', 'IE,BW', 'B&W')"""
    key = re.sub(r'[^A-Z0-9]', '', str(route).upper()) if route is not None else ''
    return key.startswith('INTEREXP') or key in ('IEBW', 'BW', 'IE')


# ---------------- น้ำหนัก + Dry Ice (For BKK) ----------------
BKK_WEIGHT_BAD_UNIT = 'น้ำหนักต่อชิ้นใน Master ผิดปกติ (ไม่นับเข้ายอด)'


def _bkk_num(v):
    try:
        f = float(v)
        return f if f == f else None
    except (TypeError, ValueError):
        return None


def _bkk_size_kg(s):
    """SIZE ใน Product Master -> kg ต่อชิ้น: '425 g' -> 0.425, '2.5 kg' -> 2.5, '900 ml' -> 0.9"""
    if not isinstance(s, str):
        return None
    m = re.search(r'([\d.]+)\s*(kg|g|gr|ml|l|lt|ltr|cl)\b', s.lower())
    if not m:
        return None
    try:
        v = float(m.group(1))
    except ValueError:
        return None
    u = m.group(2)
    return v if u in ('kg', 'l', 'lt', 'ltr') else v / 1000 if u in ('g', 'gr', 'ml') else v / 100


def bkk_line_weight(qty, unit, prod):
    """น้ำหนัก 1 รายการสินค้า -> (kg หรือ None, สถานะ)
    PCS   : Qty x Unit Net Wt (ถ้า Master ไม่มี Unit Net Wt ใช้ SIZE แทน)
    KG    : Qty คือ kg อยู่แล้ว
    CTN.. : Qty x UNIT per PACK x SIZE (CTN060 / CTN080 = ลังละ 60 / 80 ชิ้น)
    1KG1P : 1 แพ็ค = 1 kg
    น้ำหนักต่อชิ้นที่ต่างจาก SIZE เกิน 2.5 เท่า (เช่น 990 kg/ชิ้น) ไม่นับเข้ายอด กันยอดรวมเพี้ยน"""
    q = _bkk_num(qty)
    u = str(unit).strip().upper() if unit is not None else ''
    if q is None:
        return None, 'Qty ไม่ใช่ตัวเลข'
    if u == 'KG':
        return q, 'คำนวณได้'
    if prod is None:
        return None, 'ไม่พบ Product Code ใน Product Master'
    size = _bkk_size_kg(prod.get('SIZE'))
    if u == '1KG1P':
        return q * 1.0, 'คำนวณได้ (1KG1P = 1 kg)'
    if re.fullmatch(r'CTN\d*', u):
        n = _bkk_num(prod.get('UNIT per PACK'))
        if n and size:
            return q * n * size, 'คำนวณได้ (CTN: ชิ้นต่อลัง x SIZE)'
        return None, 'CTN: Master ไม่มี UNIT per PACK หรือ SIZE'
    if u == 'PCS':
        uw = _bkk_num(prod.get('Unit Net Wt (kg)'))
        if uw is None or uw <= 0:
            if size:
                return q * size, 'ประมาณจาก SIZE (Master ไม่มี Unit Net Wt)'
            return None, 'Master ไม่มี Unit Net Wt และ SIZE'
        if size and (uw / size > 2.5 or uw / size < 0.4):
            return None, BKK_WEIGHT_BAD_UNIT
        return q * uw, 'คำนวณได้'
    return None, f'หน่วย "{unit}" ยังไม่รองรับ'


def bkk_weigh_lines(lines_df, product_df):
    """เติม kg / สถานะ / กลุ่ม Dry Ice ให้ทุกรายการสินค้า (product_df=None -> ไม่คิดน้ำหนัก แต่ยังนับ 4AOB/4HFF ได้)"""
    lines = lines_df.copy()
    lines['Qty'] = pd.to_numeric(lines['Qty'], errors='coerce')
    code = lines['Product Code'].astype(str).str.strip().str.upper()
    grp = pd.Series([None] * len(lines), index=lines.index, dtype=object)
    for prefix, _ in BKK_DRYICE_GROUPS:
        grp[code.str.startswith(prefix)] = prefix
    lines['กลุ่ม'] = grp
    if product_df is None:
        lines['kg'] = None
        lines['kg_raw'] = None
        lines['สถานะ'] = 'ไม่ได้โหลด Product Master'
        lines['in_master'] = None
        return lines
    pm = {}
    for rec in product_df.to_dict('records'):
        k = str(rec.get('SKUCode', '')).strip().upper()
        if k and k not in pm:
            pm[k] = rec
    kgs, raws, sts = [], [], []
    for c, q, u in zip(code, lines['Qty'], lines['Units']):
        kg, st_ = bkk_line_weight(q, u, pm.get(c))
        raw = None
        if st_ == BKK_WEIGHT_BAD_UNIT:  # เก็บค่าที่คำนวณได้ไว้โชว์ในรายการตรวจ แต่ไม่นับเข้ายอด
            uw = _bkk_num(pm[c].get('Unit Net Wt (kg)'))
            raw = q * uw if (uw is not None and q == q) else None
        kgs.append(kg)
        raws.append(kg if kg is not None else raw)
        sts.append(st_)
    lines['kg'] = pd.to_numeric(pd.Series(kgs, index=lines.index), errors='coerce')
    lines['kg_raw'] = pd.to_numeric(pd.Series(raws, index=lines.index), errors='coerce')
    lines['สถานะ'] = sts
    lines['in_master'] = [c in pm for c in code]
    return lines


def _bkk_show(df):
    """Route มีทั้งเลขและตัวหนังสือ (1, 2, HH, PTY) -> แปลงเป็นข้อความก่อนโชว์ในตาราง"""
    out = df.copy()
    for c in ('Route', 'OrderNo'):
        if c in out.columns:
            out[c] = out[c].astype(str)
    return out


def bkk_summaries(sheets, lines):
    """sheets = [(ชื่อชีท, df), ...] ของใบงาน For BKK, lines = ผลจาก bkk_weigh_lines (เฉพาะออเดอร์ที่อยู่ในใบงาน)
    Returns dict ของตาราง: dry_route, dry_lines, weight_route, weight_sheet, bad_lines"""
    route_of, sheet_of = {}, {}
    for name, df in sheets:
        for o, r in zip(df['OrderNo'], df['Route']):
            route_of[o] = r
            sheet_of[o] = name
    lines = lines[lines['OrderNo'].isin(route_of)].copy()
    lines['Route'] = lines['OrderNo'].map(route_of)
    lines['ชีท'] = lines['OrderNo'].map(sheet_of)
    routes = sorted(set(route_of.values()), key=_bkk_route_sort_key)
    has_weight = lines['kg_raw'].notna().any() or lines['สถานะ'].ne('ไม่ได้โหลด Product Master').any()

    dry = lines[lines['กลุ่ม'].notna()]
    dry_rows = []
    for rt in routes:
        g = dry[dry['Route'] == rt]
        if g.empty:
            continue
        vals = [float(g.loc[g['กลุ่ม'] == p, 'Qty'].sum()) for p, _ in BKK_DRYICE_GROUPS]
        dry_rows.append([rt, int(g['OrderNo'].nunique())] + vals + [sum(vals)])
    dry_cols = ['Route', 'จำนวนออเดอร์'] + [f'{p} (Qty)' for p, _ in BKK_DRYICE_GROUPS] + ['รวม (Qty)']
    dry_route = pd.DataFrame(dry_rows, columns=dry_cols)
    if len(dry_route):
        tot = ['รวมทั้งหมด', int(dry['OrderNo'].nunique())] + [float(dry_route[c].sum()) for c in dry_cols[2:]]
        dry_route = pd.concat([dry_route, pd.DataFrame([tot], columns=dry_cols)], ignore_index=True)
    dry_lines = dry.assign(_rk=dry['Route'].map(_bkk_route_sort_key)).sort_values(['_rk', 'OrderNo'], kind='stable')[
        ['Route', 'OrderNo', 'ชีท', 'กลุ่ม', 'Product Code', 'Qty', 'Units']].reset_index(drop=True)

    w_rows = []
    for rt in routes:
        g = lines[lines['Route'] == rt]
        miss = int(g['kg'].isna().sum())
        w_rows.append([rt, int(g['OrderNo'].nunique()), round(float(g['kg'].sum()), 2), len(g), miss,
                       'ครบ' if miss == 0 else f'ขาด {miss} รายการ'])
    miss = int(lines['kg'].isna().sum())
    w_rows.append(['รวมทั้งหมด', int(lines['OrderNo'].nunique()), round(float(lines['kg'].sum()), 2), len(lines), miss,
                   'ครบ' if miss == 0 else f'ขาด {miss} รายการ'])
    weight_route = pd.DataFrame(w_rows, columns=['Route', 'จำนวนออเดอร์', 'น้ำหนักรวม (kg)', 'จำนวนรายการ', 'คิดไม่ได้', 'สถานะ'])

    s_rows = []
    for name, df in sheets:
        if len(df):
            g = lines[lines['ชีท'] == name]
            s_rows.append([name, len(df), round(float(g['kg'].sum()), 2)])
    s_rows.append(['รวมทั้งหมด', sum(r[1] for r in s_rows), round(float(lines['kg'].sum()), 2)])
    weight_sheet = pd.DataFrame(s_rows, columns=['ชีท', 'จำนวนออเดอร์', 'น้ำหนักรวม (kg)'])

    bad = lines[lines['kg'].isna()]
    bad_lines = bad.assign(_rk=bad['Route'].map(_bkk_route_sort_key)).sort_values(['_rk', 'OrderNo'], kind='stable')[
        ['Route', 'OrderNo', 'Product Code', 'Qty', 'Units', 'kg_raw', 'สถานะ']].rename(
        columns={'kg_raw': 'kg ที่คำนวณได้ (ไม่นับ)', 'สถานะ': 'สาเหตุ'}).reset_index(drop=True)
    # Product Code ที่ไม่มีใน Product Master (1 แถวต่อ 1 รหัส) ไว้ดาวน์โหลดไปเพิ่มใน Master
    nf = lines[lines['in_master'] == False].copy()  # noqa: E712 (None = ไม่ได้โหลด Master)
    mc_rows = []
    if len(nf):
        nf['_code'] = nf['Product Code'].astype(str).str.strip().str.upper()
        nf['_rk'] = nf['Route'].map(_bkk_route_sort_key)
        for code_, g in nf.sort_values(['_rk', 'OrderNo'], kind='stable').groupby('_code', sort=False):
            mc_rows.append({
                'Product Code': g['Product Code'].iloc[0],
                'Units': ', '.join(dict.fromkeys(str(u) for u in g['Units'] if not _cell_blank(u))),
                'จำนวนรายการ': len(g),
                'Qty รวม': float(g['Qty'].sum()),
                'จำนวนออเดอร์': int(g['OrderNo'].nunique()),
                'OrderNo': ', '.join(dict.fromkeys(str(o) for o in g['OrderNo'])),
                'Route': ', '.join(dict.fromkeys(str(r) for r in g['Route'])),
                'พบเมื่อ': datetime.now().strftime('%Y-%m-%d'),
            })
    missing_codes = pd.DataFrame(mc_rows, columns=['Product Code', 'Units', 'จำนวนรายการ', 'Qty รวม', 'จำนวนออเดอร์',
                                                   'OrderNo', 'Route', 'พบเมื่อ'])
    return {'dry_route': dry_route, 'dry_lines': dry_lines, 'weight_route': weight_route,
            'weight_sheet': weight_sheet, 'bad_lines': bad_lines, 'has_weight': bool(has_weight),
            'missing_orders': set(bad['OrderNo']), 'missing_codes': missing_codes}


def bkk_missing_codes_to_excel_bytes(df):
    """ไฟล์ Product Code ที่ไม่เจอใน Product Master"""
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        df.to_excel(writer, sheet_name='Product Code ไม่เจอ', index=False)
        ws = writer.sheets['Product Code ไม่เจอ']
        for i, w in enumerate([16, 10, 13, 10, 14, 40, 22, 12], start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
        for cell in ws[1]:
            cell.font = Font(bold=True)
        ws.freeze_panes = 'A2'
    return output.getvalue()


def process_bkk_orders(orders, info, master_df, channel_map, dates=None, service_minutes=15, retail_service_minutes=45,
                       product_df=None):
    """แยก Retail / ไม่ใช่ Retail + เติม Lat/Lon (จับด้วย ShipTo Name แบบเดียวกับแท็บประมวลผลไฟล์)
    Retail ดูจาก Channel ใน Master (คอลัมน์ C) ตาม ShipTo Name; ถ้าไม่เจอชื่อใน Master เลย
    ใช้ Cust Type ในไฟล์แทน (R = Retail) แล้วนับไว้รายงาน
    ลำดับการแยกชีท:
      1) Route เป็น InterExp / IE,B&W -> ชีท InterExp-B&W เสมอ (ไม่ว่าจะเป็น Retail หรือไม่)
      2) ช่อง Cust Type ในไฟล์ว่าง -> ชีท No Cust Type (เก็บไว้ให้คนตรวจ ไม่เดาว่าเป็นประเภทไหน)
      3) ที่เหลือแบ่ง Retail / ไม่ใช่ Retail
    Returns (normal_df, retail_df, carrier_df, notype_df, info)"""
    df = orders.copy()
    if dates is not None:
        keep = set(str(d) for d in dates)
        df = df[df['Order Date'].astype(str).isin(keep)]

    by_ship = master_df.drop_duplicates('norm_ship').set_index('norm_ship')
    name_col = 'ShipTo Name' if info.get('has_shipto_col', True) else 'Customer Name'
    multi_pin_keys = master_multi_pin_names(master_df)

    forced_keys = {_bkk_name_key(n) for n in BKK_FORCE_FOOD_SERVICE}
    lats, lons, is_retail, cust_types, by_cust_type, forced_fs = [], [], [], [], 0, 0
    for _, r in df.iterrows():
        key = normalize(r[name_col])
        no_coord_key = False
        if not key and name_col == 'ShipTo Name':
            # ShipTo Name ว่างเฉพาะออเดอร์นี้ -> ใช้ Customer Name (ถ้าชื่อนี้มีหลายหมุดใน Master จะไม่เดาพิกัด)
            key = normalize(r['Customer Name'])
            no_coord_key = key in multi_pin_keys
        ch = channel_map.get(key) if key else None
        if _bkk_name_key(r['ShipTo Name']) in forced_keys or _bkk_name_key(r['Customer Name']) in forced_keys:
            # ข้อยกเว้นที่กำหนดไว้: ถือเป็น Food Service เสมอ และเปลี่ยน Cust Type เป็น F
            forced_fs += 1
            is_retail.append(False)
            cust_types.append('F')
        elif ch is not None:
            is_retail.append(ch.strip().lower() == 'retail')
            cust_types.append(r['Cust Type'])
        else:
            by_cust_type += 1
            is_retail.append(str(r['Cust Type']).strip().upper() == 'R')
            cust_types.append(r['Cust Type'])
        if key and not no_coord_key and key in by_ship.index:
            hit = by_ship.loc[key]
            lats.append(hit['Latitude'])
            lons.append(hit['Longitude'])
        else:
            lats.append(None)
            lons.append(None)

    df['Cust Type'] = cust_types
    df['Remark'] = None
    df['เวลาเข้า'] = None
    df['เวลาออก'] = None
    df[BKK_SERVICE_COL] = None
    # น้ำหนักต่อออเดอร์ = ผลรวม kg ทุกรายการสินค้าของออเดอร์นั้นใน Sheet1 / 4AOB 4HFF = ผลรวม Qty (1 = 1 ไม่หาร)
    lines_df = info.get('lines_df')
    if lines_df is None:
        lines_df = pd.DataFrame(columns=['OrderNo', 'Product Code', 'Qty', 'Units'])
    weighed = bkk_weigh_lines(lines_df[lines_df['OrderNo'].isin(set(df['OrderNo']))], product_df)
    by_order = weighed.groupby('OrderNo')
    df[BKK_WEIGHT_COL] = df['OrderNo'].map(by_order['kg'].sum(min_count=1)).round(2) if product_df is not None else None
    for prefix, col_name in BKK_DRYICE_GROUPS:
        g = weighed[weighed['กลุ่ม'] == prefix].groupby('OrderNo')['Qty'].sum()
        df[col_name] = df['OrderNo'].map(g)
    df['Latitude'] = lats
    df['Longitude'] = lons
    df['_retail'] = is_retail
    df['_rk'] = df['Route'].map(_bkk_route_sort_key)
    df['_ck'] = df['Customer Name'].map(lambda v: '' if v is None else str(v))
    df['_ok'] = df['OrderNo'].map(lambda v: (0, v, '') if isinstance(v, int) else (1, 0, str(v)))
    df = df.sort_values(['_rk', '_ck', '_ok'], kind='stable')

    df['_carrier'] = df['Route'].map(_bkk_is_carrier_route)
    df['_notype'] = df['Cust Type'].map(_cell_blank) & ~df['_carrier']
    rest = ~df['_carrier'] & ~df['_notype']
    # Service Time: เวลามาตรฐานต่อร้าน (Retail 45 นาที / อื่นๆ 15 นาที) หารด้วยจำนวน SO ที่
    # ShipTo Name เดียวกันในชีทเดียวกัน เช่น 3 SO ไปร้านเดียวกัน 15 นาที -> SO ละ 5 นาที
    if len(df):
        bucket = df['_carrier'].map({True: 'C', False: ''}) + df['_notype'].map({True: 'N', False: ''}) \
            + (rest & df['_retail']).map({True: 'R', False: ''})
        ship = df['ShipTo Name'].map(lambda v: normalize(v) or None)
        ship = ship.fillna(pd.Series([f'__row{i}' for i in range(len(df))], index=df.index))
        key = bucket + '|' + ship
        standard = df['_retail'].map({True: float(retail_service_minutes), False: float(service_minutes)})
        df[BKK_SERVICE_COL] = (standard / key.map(key.value_counts())).round(2)
    normal = df[rest & ~df['_retail']][BKK_COLUMNS].reset_index(drop=True)
    retail = df[rest & df['_retail']][BKK_COLUMNS].reset_index(drop=True)
    carrier = df[df['_carrier']][BKK_COLUMNS].reset_index(drop=True)
    notype = df[df['_notype']][BKK_COLUMNS].reset_index(drop=True)
    # คอลัมน์ Route ในไฟล์ผลลัพธ์ = Route จากคอลัมน์ A ของ Sheet1 (ออเดอร์ที่มีหลาย Route ใช้ Route แรกที่เจอ)
    out_info = dict(info)
    out_info.update({
        'selected_orders': len(df), 'normal': len(normal), 'retail': len(retail), 'carrier': len(carrier),
        'notype': len(notype), 'coords_notype': int(notype['Latitude'].notna().sum()),
        'carrier_notype': int((df['_carrier'] & df['Cust Type'].map(_cell_blank)).sum()),
        'carrier_retail': int((df['_carrier'] & df['_retail']).sum()),
        'retail_by_cust_type': by_cust_type, 'forced_fs': forced_fs, 'service_minutes': service_minutes, 'retail_service_minutes': retail_service_minutes,
        'coords_normal': int(normal['Latitude'].notna().sum()), 'coords_retail': int(retail['Latitude'].notna().sum()),
        'coords_carrier': int(carrier['Latitude'].notna().sum()),
        'weighed_lines': weighed,
    })
    out_info.pop('lines_df', None)
    return normal, retail, carrier, notype, out_info


def bkk_sheet1_clean(file_bytes):
    """Sheet1 ต้นฉบับ (ทุกวันที่ ไม่กรองอะไร) ตัดเฉพาะแถวที่ว่างทั้งแถว และคอลัมน์คั่นที่ว่างทั้งคอลัมน์
    ค่าในทุกช่องที่มีข้อมูลคงเดิม"""
    xl = pd.ExcelFile(io.BytesIO(file_bytes))
    sheet = 'Sheet1' if 'Sheet1' in xl.sheet_names else xl.sheet_names[0]
    raw = xl.parse(sheet, header=None)
    blank = raw.apply(lambda col: col.map(_cell_blank))
    return raw.loc[~blank.all(axis=1), ~blank.all(axis=0)].reset_index(drop=True)


def _bkk_write_summary_sheets(wb, summary):
    """ชีท 'สรุป Dry Ice' และ 'สรุปน้ำหนัก' ต่อท้ายใบงาน"""
    thin = Side(style='thin')
    bd = Border(top=thin, bottom=thin, left=thin, right=thin)
    fh = Font(name='Cordia New', size=14, bold=True)
    fb = Font(name='Cordia New', size=14)
    ft = Font(name='Cordia New', size=16, bold=True)

    def table(ws, top, df, bold_last=False):
        for c, h in enumerate(df.columns, 1):
            x = ws.cell(row=top, column=c, value=h)
            x.font, x.border = fh, bd
            x.alignment = Alignment(horizontal='center')
        for i, rec in enumerate(df.itertuples(index=False), top + 1):
            for c, v in enumerate(rec, 1):
                if v is not None and isinstance(v, float) and pd.isna(v):
                    v = None
                x = ws.cell(row=i, column=c, value=v)
                x.font, x.border = fb, bd
                if isinstance(v, float):
                    x.number_format = '#,##0.00' if v != int(v) else '#,##0'
        if bold_last and len(df):
            for c in range(1, len(df.columns) + 1):
                ws.cell(row=top + len(df), column=c).font = fh
        return top + 1 + len(df)

    ws = wb.create_sheet(title='สรุป Dry Ice')
    ws['A1'] = 'สรุปสินค้า 4AOB / 4HFF (สำหรับเตรียม Dry Ice)'
    ws['A1'].font = ft
    if len(summary['dry_route']):
        ws['A3'] = 'ต่อ Route  (Qty 1 = 1 ไม่หาร)'
        ws['A3'].font = fh
        end = table(ws, 4, summary['dry_route'], bold_last=True)
        ws.cell(row=end + 1, column=1, value='ต่อออเดอร์').font = fh
        table(ws, end + 2, summary['dry_lines'])
    else:
        ws['A3'] = 'ไม่มีสินค้า 4AOB / 4HFF ในไฟล์นี้'
        ws['A3'].font = fb
    for c, wd in enumerate([16, 14, 16, 14, 14, 14, 12], 1):
        ws.column_dimensions[get_column_letter(c)].width = wd

    if not summary.get('has_weight'):
        return
    ws = wb.create_sheet(title='สรุปน้ำหนัก')
    ws['A1'] = 'น้ำหนักต่อ Route (คำนวณจาก Sheet1 + Product Master)'
    ws['A1'].font = ft
    ws['A2'] = 'ต่อ Route'
    ws['A2'].font = fh
    end = table(ws, 3, summary['weight_route'], bold_last=True)
    ws.cell(row=end + 1, column=1, value='ต่อชีท').font = fh
    end = table(ws, end + 2, summary['weight_sheet'], bold_last=True)
    if len(summary['bad_lines']):
        ws.cell(row=end + 1, column=1, value='รายการที่คิดน้ำหนักไม่ได้ (ไม่ได้นับเข้ายอด)').font = fh
        table(ws, end + 2, summary['bad_lines'])
    for c, wd in enumerate([16, 14, 18, 14, 12, 22, 45], 1):
        ws.column_dimensions[get_column_letter(c)].width = wd


def bkk_to_excel_bytes(normal_df, retail_df, carrier_df, notype_df=None, sheet1_df=None, summary=None):
    """ชีท Order / Retail + InterExp-B&W (+ No Cust Type ถ้ามีออเดอร์ที่ช่อง Cust Type ว่าง)
    หน้าตาเดียวกับใบงาน Sheet2: หัวกระดาษ ทะเบียนรถ/วันที่/เลขไมล์,
    หัวตารางแถว 4, ฟอนต์ Cordia New, เส้นขอบทุกช่อง + Lat/Lon ต่อท้าย"""
    wb = Workbook()
    wb.remove(wb.active)
    thin = Side(style='thin')
    border = Border(top=thin, bottom=thin, left=thin, right=thin)
    f_top = Font(name='Cordia New', size=12, bold=True)
    f_head = Font(name='Cordia New', size=14, bold=True, underline='single')
    f_body = Font(name='Cordia New', size=14)
    center_cols = {'Cust Type', 'Route'}
    f_bold = Font(name='Cordia New', size=14, bold=True)
    fill_warn = PatternFill('solid', fgColor='FFE699')  # น้ำหนักไม่ครบ (มีรายการที่คิดไม่ได้)
    fill_ice = PatternFill('solid', fgColor='DDEBF7')   # ออเดอร์ที่มี 4AOB / 4HFF
    missing_orders = summary['missing_orders'] if summary and summary.get('has_weight') else set()
    w_idx = BKK_COLUMNS.index(BKK_WEIGHT_COL) + 1
    ice_idx = [BKK_COLUMNS.index(c) + 1 for _, c in BKK_DRYICE_GROUPS]

    # Sheet1 ต้นฉบับ (ตัดแถว/คอลัมน์ว่างแล้ว) เก็บไว้เป็นชีทแรก
    if sheet1_df is not None and len(sheet1_df):
        ws1 = wb.create_sheet(title='Sheet1')
        f_s1 = Font(name='Cordia New', size=14)
        f_s1_head = Font(name='Cordia New', size=14, bold=True)
        for r_i, row in enumerate(sheet1_df.itertuples(index=False), start=1):
            for c_i, v in enumerate(row, start=1):
                if _cell_blank(v):
                    continue
                cell = ws1.cell(row=r_i, column=c_i, value=v)
                cell.font = f_s1_head if r_i == 1 else f_s1
        for c_i, col in enumerate(sheet1_df.columns, start=1):
            width = max((len(str(v)) for v in sheet1_df[col] if not _cell_blank(v)), default=6)
            ws1.column_dimensions[get_column_letter(c_i)].width = min(max(width + 2, 8), 50)

    sheets = bkk_output_sheets(normal_df, retail_df, carrier_df, notype_df)
    for title, df in sheets:
        ws = wb.create_sheet(title=title)
        ws['B1'] = 'ทะเบียนรถ____________________'
        ws['F1'] = 'วันที่______________________'
        ws['B2'] = 'เลขไมล์ออกจากคลัง___________________________'
        ws['F2'] = 'เลขไมล์กลับถึงคลัง___________________________'
        for ref in ('B1', 'F1', 'B2', 'F2'):
            ws[ref].font = f_top
        for r in (1, 2, 3):
            ws.row_dimensions[r].height = 18.75

        for c, name in enumerate(BKK_COLUMNS, start=1):
            cell = ws.cell(row=4, column=c, value=name)
            cell.font = f_head
            cell.border = border
            cell.alignment = Alignment(horizontal='center', vertical='center')
        ws.row_dimensions[4].height = 21

        for i, rec in enumerate(df.itertuples(index=False), start=5):
            for c, (name, v) in enumerate(zip(BKK_COLUMNS, rec), start=1):
                if v is not None and isinstance(v, float) and pd.isna(v):
                    v = None
                cell = ws.cell(row=i, column=c, value=v)
                cell.font = f_body
                cell.border = border
                cell.alignment = Alignment(horizontal='center' if name in center_cols else 'left', vertical='center')
                if name == BKK_WEIGHT_COL:
                    cell.number_format = '#,##0.00'
                    if rec.OrderNo in missing_orders:
                        cell.fill = fill_warn
                elif c in ice_idx and v is not None:
                    cell.fill = fill_ice
            ws.row_dimensions[i].height = 21.75

        # แถวรวมใต้คอลัมน์น้ำหนัก / 4AOB / 4HFF
        if len(df):
            tr = 5 + len(df)
            ws.cell(row=tr, column=w_idx - 1, value='รวม').font = f_bold
            for c in [w_idx] + ice_idx:
                L = get_column_letter(c)
                x = ws.cell(row=tr, column=c, value=f'=SUM({L}5:{L}{tr - 1})')
                x.font = f_bold
                x.number_format = '#,##0.00' if c == w_idx else 'General'

        # ท้ายใบงาน: เว้น 2 แถว แล้วช่องลงชื่อคนขับ + หมายเหตุถึงคนขับ
        foot = 5 + len(df) + 2
        ws.cell(row=foot, column=2, value=BKK_FOOTER_SIGN[0]).font = f_body
        ws.cell(row=foot, column=6, value=BKK_FOOTER_SIGN[1]).font = f_body
        ws.row_dimensions[foot].height = 21
        for k, note in enumerate(BKK_FOOTER_NOTES, start=1):
            ws.cell(row=foot + k, column=2, value=note).font = f_body
            ws.row_dimensions[foot + k].height = 21

        for c, w in enumerate(BKK_COL_WIDTHS, start=1):
            ws.column_dimensions[get_column_letter(c)].width = w
        ws.print_title_rows = '1:4'
        ws.page_setup.orientation = 'landscape'
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True

    # สรุป Dry Ice / น้ำหนัก แสดงบนหน้าแอพอย่างเดียว ไม่ใส่เป็นชีทในไฟล์ (ตามที่ผู้ใช้ขอ)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _render_bkk():
    if not MASTER_SHEET_CSV_URL:
        st.error("⚠️ ยังไม่ได้ตั้งค่า MASTER_SHEET_CSV_URL — ดูวิธีตั้งค่าใน DEPLOY_INSTRUCTIONS.md")
        return
    st.markdown("##### 🏙️ For BKK — ทำใบงานจาก Sheet1")

    bkk_file = st.file_uploader("อัพโหลดไฟล์ route", type=['xls', 'xlsx'], label_visibility="collapsed", key="bkk_uploader")

    if bkk_file is not None:
        bkk_bytes = bkk_file.read()
        try:
            bkk_orders, bkk_info = read_bkk_orders(bkk_bytes)
        except Exception as e:
            st.error(f"อ่านไฟล์ไม่ได้: {e}")
            return

        bkk_dates = sorted({str(d) for d in bkk_orders['Order Date'] if d is not None})
        st.markdown(f"""
        <div class="rcm-card">
            📄 ชีท <b>{bkk_info['sheet']}</b>: {bkk_info['lines']:,} รายการสินค้า → <b>{bkk_info['orders']:,} ออเดอร์</b> (OrderNo ไม่ซ้ำ)
        </div>
        """, unsafe_allow_html=True)

        if len(bkk_dates) > 1:
            bkk_sel_dates = st.multiselect(
                "Order Date ที่จะเอา (ค่าเริ่มต้น = ทุกวันที่มีใน Sheet1)",
                options=bkk_dates, default=bkk_dates, key="bkk_dates",
            )
        else:
            bkk_sel_dates = bkk_dates

        if not bkk_sel_dates:
            st.warning("⚠️ กรุณาเลือก Order Date อย่างน้อย 1 วัน")
            return

        st.caption("⏱️ เวลาบริการมาตรฐานต่อร้าน (นาที) — หารเท่า ๆ กันตามจำนวน SO ที่ ShipTo Name เดียวกัน")
        svc_c1, svc_c2 = st.columns(2)
        with svc_c1:
            bkk_service = st.number_input("Food Service", min_value=1.0, max_value=240.0, value=15.0, step=1.0, key="bkk_service_min")
        with svc_c2:
            bkk_service_retail = st.number_input("Retail", min_value=1.0, max_value=240.0, value=45.0, step=1.0, key="bkk_service_min_retail")
        bkk_clicked = st.button("🚀 ทำใบงาน For BKK", type="primary", use_container_width=True, key="bkk_btn")

        if bkk_clicked:
            with st.spinner("🛰️ กำลังแยก Retail เติมพิกัด และคิดน้ำหนัก..."):
                try:
                    master_df = load_master_data(MASTER_SHEET_CSV_URL)
                    bkk_channels = load_master_channels(MASTER_SHEET_CSV_URL)
                    bkk_product, bkk_product_err = None, None
                    if PRODUCT_SHEET_CSV_URL:
                        try:
                            bkk_product = load_product_master(PRODUCT_SHEET_CSV_URL)
                        except Exception as pe:
                            bkk_product_err = str(pe)
                    else:
                        bkk_product_err = 'ยังไม่ได้ตั้งค่า PRODUCT_SHEET_CSV_URL'
                    bkk_normal, bkk_retail, bkk_carrier, bkk_notype, bkk_result = process_bkk_orders(
                        bkk_orders, bkk_info, master_df, bkk_channels,
                        dates=bkk_sel_dates if len(bkk_dates) > 1 else None,
                        service_minutes=float(bkk_service),
                        retail_service_minutes=float(bkk_service_retail),
                        product_df=bkk_product,
                    )
                    bkk_sheets = bkk_output_sheets(bkk_normal, bkk_retail, bkk_carrier, bkk_notype)
                    bkk_summary = bkk_summaries(bkk_sheets, bkk_result['weighed_lines'])
                    st.session_state['bkk_xlsx'] = bkk_to_excel_bytes(bkk_normal, bkk_retail, bkk_carrier, bkk_notype,
                                                                      sheet1_df=bkk_sheet1_clean(bkk_bytes),
                                                                      summary=bkk_summary)
                    st.session_state['bkk_summary'] = bkk_summary
                    st.session_state['bkk_product_err'] = bkk_product_err
                    st.session_state['bkk_result'] = bkk_result
                    st.session_state['bkk_filename'] = bkk_file.name
                except Exception as e:
                    st.error(f"เกิดข้อผิดพลาด: {e}")
                    st.exception(e)

        if 'bkk_result' in st.session_state and st.session_state.get('bkk_filename') == bkk_file.name:
            res = st.session_state['bkk_result']
            st.markdown(f"""
            <div class="rcm-card rcm-card-success">
                🎯 <b>เสร็จแล้ว!</b> {res['selected_orders']:,} ออเดอร์ → Order {res['normal']:,} + Retail {res['retail']:,} + InterExp/B&W {res['carrier']:,} + ไม่มี Cust Type {res['notype']:,}
            </div>
            """, unsafe_allow_html=True)

            bc1, bc2, bc4 = st.columns(3)
            rc_total = res['retail'] + res['carrier']
            with bc1:
                stat_box("SHEET: Order", f"{res['normal']:,}",
                          f"มีพิกัด {res['coords_normal']:,} / {res['normal']:,}")
            with bc2:
                stat_box(f"SHEET: {BKK_RETAIL_CARRIER_SHEET}", f"{rc_total:,}",
                          f"Retail {res['retail']:,} · InterExp-B&W {res['carrier']:,} · "
                          f"มีพิกัด {res['coords_retail'] + res['coords_carrier']:,} / {rc_total:,}")
            with bc4:
                stat_box("SHEET: No Cust Type", f"{res['notype']:,}",
                          f"มีพิกัด {res['coords_notype']:,} / {res['notype']:,}" if res['notype'] else "ไม่มี — ไม่สร้างชีทนี้")

            bkk_notes = []
            if res['retail_by_cust_type']:
                bkk_notes.append(f"{res['retail_by_cust_type']:,} ออเดอร์ไม่เจอ ShipTo Name ใน Master — แยก Retail จาก Cust Type ในไฟล์แทน (R = Retail)")
            if res['notype']:
                bkk_notes.append(f"{res['notype']:,} ออเดอร์ช่อง Cust Type ว่าง — แยกไว้ในชีท {BKK_NOTYPE_SHEET} ให้ตรวจเอง (ไม่ได้ลบ และไม่ได้เดาประเภท)")
            if res.get('carrier_notype'):
                bkk_notes.append(f"{res['carrier_notype']:,} ออเดอร์ InterExp-B&W (ชีท {BKK_RETAIL_CARRIER_SHEET}) ช่อง Cust Type ว่าง (อยู่ชีทนั้นตาม Route)")
            if res.get('forced_fs'):
                bkk_notes.append(f"{res['forced_fs']:,} ออเดอร์ของ {', '.join(BKK_FORCE_FOOD_SERVICE)} ถูกเปลี่ยนเป็น F (Food Service) ตามที่ตั้งไว้")
            if res['multi_route']:
                mr = ', '.join(f"{k} ({' / '.join(str(x) for x in v)})" for k, v in list(res['multi_route'].items())[:10])
                bkk_notes.append(f"{len(res['multi_route']):,} ออเดอร์มีมากกว่า 1 Route ใน Sheet1 — ใช้ Route แรกที่เจอในการแยกชีท: {mr}")
            if res['no_order_rows']:
                bkk_notes.append(f"{res['no_order_rows']:,} แถวมีข้อมูลแต่ไม่มี OrderNo — ไม่ได้เอามา")
            if bkk_notes:
                st.markdown('<div class="rcm-card rcm-card-warn">⚠️ <b>ข้อสังเกต</b><br>' + '<br>'.join('• ' + n for n in bkk_notes) + '</div>',
                            unsafe_allow_html=True)

            summ = st.session_state.get('bkk_summary')
            if summ is not None:
                st.write("")
                st.markdown("##### 🧊 Dry Ice — สินค้า 4AOB / 4HFF")
                if len(summ['dry_route']):
                    tot = summ['dry_route'].iloc[-1]
                    dc = st.columns(len(BKK_DRYICE_GROUPS) + 2)
                    with dc[0]:
                        stat_box("ออเดอร์ที่ต้องใช้ Dry Ice", f"{int(tot['จำนวนออเดอร์']):,}")
                    for k, (prefix, _) in enumerate(BKK_DRYICE_GROUPS, start=1):
                        with dc[k]:
                            stat_box(f"{prefix} (Qty)", f"{tot[f'{prefix} (Qty)']:,.0f}")
                    with dc[-1]:
                        stat_box("รวม (Qty)", f"{tot['รวม (Qty)']:,.0f}", "Qty 1 = 1")
                    st.dataframe(_bkk_show(summ['dry_route']), use_container_width=True, hide_index=True)
                    with st.expander(f"ดูรายการ 4AOB / 4HFF ทั้งหมด ({len(summ['dry_lines']):,} รายการ)"):
                        st.dataframe(_bkk_show(summ['dry_lines']), use_container_width=True, hide_index=True)
                else:
                    st.markdown('<div class="rcm-card">ไม่มีสินค้า 4AOB / 4HFF ในไฟล์นี้</div>', unsafe_allow_html=True)

                st.write("")
                st.markdown("##### ⚖️ น้ำหนัก (รวมจากทุกรายการสินค้าใน Sheet1)")
                if summ['has_weight']:
                    wt = summ['weight_route'].iloc[-1]
                    wc_ = st.columns(len(summ['weight_sheet']))
                    for k, rec in enumerate(summ['weight_sheet'].itertuples(index=False)):
                        with wc_[k]:
                            stat_box(rec[0], f"{rec[2]:,.2f} kg", f"{rec[1]:,} ออเดอร์")
                    if wt['คิดไม่ได้']:
                        n_all, n_bad = int(wt['จำนวนรายการ']), int(wt['คิดไม่ได้'])
                        st.markdown(
                            f'<div class="rcm-card rcm-card-warn">⚠️ คิดน้ำหนักได้ {n_all - n_bad:,} / {n_all:,} รายการ — '
                            f'อีก {n_bad:,} รายการไม่ได้นับเข้ายอด (ยอดจริงจะมากกว่านี้) ดูสาเหตุด้านล่าง'
                            ' · ในไฟล์ ช่องน้ำหนักสีเหลือง = ออเดอร์ที่น้ำหนักยังไม่ครบ</div>',
                            unsafe_allow_html=True)
                    st.dataframe(_bkk_show(summ['weight_route']), use_container_width=True, hide_index=True)
                    if len(summ['bad_lines']):
                        with st.expander(f"รายการที่คิดน้ำหนักไม่ได้ ({len(summ['bad_lines']):,}) — ส่วนใหญ่ต้องไปเพิ่ม/แก้ใน Product Master"):
                            st.dataframe(_bkk_show(summ['bad_lines']), use_container_width=True, hide_index=True)
                    if len(summ.get('missing_codes', [])):
                        mc = summ['missing_codes']
                        st.download_button(
                            f"📋 ดาวน์โหลด Product Code ที่ไม่เจอใน Product Master ({len(mc):,} รหัส)",
                            data=bkk_missing_codes_to_excel_bytes(mc),
                            file_name=f"product_code_not_found_{bkk_file.name.rsplit('.', 1)[0]}.xlsx",
                            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                            use_container_width=True,
                            key="bkk_missing_codes_btn",
                            help="รหัสสินค้าใน Sheet1 ที่ไม่มีใน Product Master — เอาไปเพิ่มใน Google Sheets แล้วกดทำใบงานใหม่ น้ำหนักจะครบขึ้น",
                        )
                else:
                    st.markdown(
                        f'<div class="rcm-card rcm-card-warn">⚠️ ไม่ได้คิดน้ำหนัก — โหลด Product Master ไม่ได้ '
                        f'({st.session_state.get("bkk_product_err") or "-"}) · ใบงานยังทำได้ตามปกติ แค่ช่องน้ำหนักว่าง</div>',
                        unsafe_allow_html=True)

            st.write("")
            st.download_button(
                "⬇️ ดาวน์โหลดใบงาน For BKK (Excel: Sheet1 + Order + Retail + InterExp-B&W + No Cust Type)",
                data=st.session_state['bkk_xlsx'],
                file_name=f"{bkk_file.name.rsplit('.', 1)[0]}_For_BKK.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                use_container_width=True,
                key="bkk_download_btn",
            )


# ============================================================
# TAB: FOR PKT — ไฟล์ดิบภูเก็ต (รายการ Invoice) -> ใบงาน 1 แถวต่อ 1 SO + Lat/Lon + Service Time
#   ไม่มี Dry Ice / น้ำหนัก / ท้ายใบงาน
# ============================================================
PKT_SERVICE_COL = 'Service Time (นาที)'
PKT_COLUMNS = ['Order Date', 'OrderNo', 'Invoice', 'Customers Code', 'Cust Type', 'Customer Name', 'ShipTo Name',
               'Warehouse', 'Sales Person', 'Route', 'Remark', 'เวลาเข้า', 'เวลาออก', PKT_SERVICE_COL,
               'Latitude', 'Longitude']
PKT_COL_WIDTHS = [11, 11, 18, 13, 9, 45, 45, 11, 15, 12, 22, 8, 8, 15, 12.5, 12.5]
PKT_NO_SO_SHEET = 'ไม่มี SO'
PKT_JOB_SHEET = 'Order'
# ชีทใบงานตามเทมเพลตที่ทีมภูเก็ตใช้ (Route เว้นว่างไว้ให้กรอกเอง)
PKT_JOB_COLUMNS = ['Inv Date', 'Order', 'Cust. Name', 'Ship To', 'Term', 'Sales Person', 'Legal Number',
                   'Total Amount ', 'Route', 'อุณหภูมิก่อนโหลด', 'อุณหภูมิเวลาเข้าส่ง', 'อุณหภูมิเวลาออกส่ง',
                   PKT_SERVICE_COL, 'Latitude', 'Longitude']
PKT_JOB_WIDTHS = [12, 10, 42.6, 42.7, 24.9, 13, 16, 13, 11, 15.4, 15.5, 16, 15, 12.5, 12.5]
PKT_JOB_TOP = ('ทะเบียนรถ___________________', 'เลขไมล์ ออกจากคลัง________________________')
PKT_JOB_FOOTER = 'รายชื่อ พขร.'
PKT_FS_GROUPS = {'FS', 'ACCO', 'MINO', 'HYAT'}   # คอลัมน์ Group ในไฟล์ดิบ: FS / ACCO(R) / Mino(r) / Hyatt -> Food Service


def _pkt_type_from_channel(ch):
    """Channel ใน Customer Master -> F / R / None"""
    if ch is None:
        return None
    c = str(ch).strip().lower()
    if c == 'retail':
        return 'R'
    if c == 'food service' or c.startswith('fs'):
        return 'F'
    return None


def _pkt_type_from_group(g):
    """คอลัมน์ Group ในไฟล์ดิบ -> F / R / None (ใช้เมื่อหา Channel ใน Master ไม่เจอ)"""
    if _cell_blank(g):
        return None
    k = str(g).strip().upper()
    if k in PKT_FS_GROUPS or k.startswith('FS'):
        return 'F'
    if k in ('RT', 'R', 'RETAIL') or k.startswith('RT'):
        return 'R'
    return None


def read_pkt_orders(file_bytes):
    """อ่านไฟล์ดิบภูเก็ต (1 แถว = 1 Invoice) -> 1 แถวต่อ 1 SO (คอลัมน์ Order)
    - ตัดแถวว่างทั้งแถว
    - SO เดียวกันหลาย Invoice -> รวมเป็นแถวเดียว (เก็บเลข Invoice ไว้ทุกใบ)
    - Invoice ที่ไม่มี SO (Order ว่างหรือ 0) -> ไม่ทิ้ง แยกไว้ต่างหาก 1 แถวต่อ 1 Invoice
    Returns (orders_df, no_so_df, info)"""
    xl = pd.ExcelFile(io.BytesIO(file_bytes))
    sheet = 'Sheet1' if 'Sheet1' in xl.sheet_names else xl.sheet_names[0]
    raw = xl.parse(sheet, header=None).values.tolist()
    header_row, c_order = find_header_row(raw, ['order', 'orderno', 'order no', 'order no.', 'so', 'so no', 'so no.'])
    if header_row is None:
        raise ValueError(f'ไม่พบคอลัมน์ Order (เลข SO) ในชีท "{sheet}"')
    header = [v.strip().lower() if isinstance(v, str) else '' for v in raw[header_row]]

    def col(*names):
        for i, h in enumerate(header):
            if h in names:
                return i
        return None

    c_date = col('inv date', 'order date', 'invoice date', 'date')
    c_inv = col('invoice', 'invoice no', 'invoice no.')
    c_code = col('cust. id', 'cust id', 'cust.id', 'cust code', 'customer code', 'customers code')
    c_cname = col('cust. name', 'cust name', 'cust.name', 'customer name')
    c_ship = col('ship to', 'shipto name', 'ship to name', 'shipto')
    c_group = col('group')
    c_wh = col('warehouse')
    c_sales = col('sales person', 'salesperson')
    c_term = col('term')
    c_legal = col('legal number', 'legal no', 'legal no.')
    c_amt = col('total amount', 'total')

    def get(row, c):
        if c is None or c >= len(row) or _cell_blank(row[c]):
            return None
        v = row[c]
        return v.strip() if isinstance(v, str) else v

    def fmt_date(v):
        if isinstance(v, (pd.Timestamp, datetime)):
            return v.strftime('%Y/%m/%d')
        return None if v is None else str(v)

    orders, seen, no_so = [], {}, []
    lines = 0
    for row in raw[header_row + 1:]:
        if all(_cell_blank(v) for v in row):
            continue
        lines += 1
        so = _bkk_clean_key(get(row, c_order))
        inv = _bkk_clean_key(get(row, c_inv))
        rec = {
            'Order Date': fmt_date(get(row, c_date)), 'OrderNo': so, 'Invoice': None if inv is None else str(inv),
            'Customers Code': get(row, c_code), 'Group': get(row, c_group), 'Customer Name': get(row, c_cname),
            'ShipTo Name': get(row, c_ship), 'Warehouse': get(row, c_wh), 'Sales Person': get(row, c_sales),
            'Term': get(row, c_term), 'Legal Number': get(row, c_legal),
            'Total Amount': get(row, c_amt), '_date': get(row, c_date),
        }
        if so in (None, 0, '0', ''):
            rec['OrderNo'] = None
            no_so.append(rec)
            continue
        if so in seen:
            first = seen[so]
            if rec['Invoice'] and rec['Invoice'] not in (first['Invoice'] or '').split(', '):
                first['Invoice'] = rec['Invoice'] if not first['Invoice'] else first['Invoice'] + ', ' + rec['Invoice']
                # SO เดียวหลาย Invoice: รวมยอดเงิน และเก็บ Legal Number ทุกใบ
                if isinstance(rec['Total Amount'], (int, float)):
                    first['Total Amount'] = (first['Total Amount'] or 0) + rec['Total Amount']
                if rec['Legal Number'] and first['Legal Number'] and rec['Legal Number'] not in first['Legal Number']:
                    first['Legal Number'] = f"{first['Legal Number']}, {rec['Legal Number']}"
            for k, v in rec.items():
                if first.get(k) is None and v is not None:
                    first[k] = v
            continue
        seen[so] = rec
        orders.append(rec)

    cols = ['Order Date', 'OrderNo', 'Invoice', 'Customers Code', 'Group', 'Customer Name', 'ShipTo Name',
            'Warehouse', 'Sales Person', 'Term', 'Legal Number', 'Total Amount', '_date']
    info = {'sheet': sheet, 'lines': lines, 'orders': len(orders), 'no_so': len(no_so),
            'multi_invoice': sum(1 for r in orders if r['Invoice'] and ', ' in r['Invoice'])}
    return pd.DataFrame(orders, columns=cols), pd.DataFrame(no_so, columns=cols), info


def process_pkt_orders(orders, no_so, master_df, channel_map, dates=None, fs_minutes=20, retail_minutes=45):
    """เติม Cust Type (Channel ใน Master ก่อน ไม่เจอค่อยดูคอลัมน์ Group), Lat/Lon และ Service Time
    Service Time: F = fs_minutes, R = retail_minutes หารตามจำนวน SO ที่ ShipTo Name เดียวกัน (ประเภทอื่นเว้นว่าง)
    ShipTo Name ว่าง -> ใช้ Customer Name แทน (ถ้าชื่อนั้นมีหลายหมุดใน Master จะไม่เดาพิกัด)
    Returns (orders_df, no_so_df, info)"""
    by_ship = master_df.drop_duplicates('norm_ship').set_index('norm_ship')
    multi_pin_keys = master_multi_pin_names(master_df)

    def fill(df):
        df = df.copy()
        if dates is not None:
            keep = set(str(d) for d in dates)
            df = df[df['Order Date'].astype(str).isin(keep)]
        lats, lons, types, keys, by_group = [], [], [], [], 0
        for _, r in df.iterrows():
            key = normalize(r['ShipTo Name'])
            no_coord = False
            if not key:
                key = normalize(r['Customer Name'])
                no_coord = key in multi_pin_keys
            keys.append(key or None)
            t = _pkt_type_from_channel(channel_map.get(key)) if key else None
            if t is None:
                t = _pkt_type_from_group(r['Group'])
                if t is not None:
                    by_group += 1
            types.append(t)
            if key and not no_coord and key in by_ship.index:
                hit = by_ship.loc[key]
                lats.append(hit['Latitude'])
                lons.append(hit['Longitude'])
            else:
                lats.append(None)
                lons.append(None)
        df['Cust Type'] = types
        df['Latitude'] = lats
        df['Longitude'] = lons
        df['_key'] = keys
        for c in ('Route', 'Remark', 'เวลาเข้า', 'เวลาออก'):
            df[c] = None
        return df, by_group

    # Invoice ที่ไม่มี SO (Order = 0) อยู่ชีทเดียวกับ SO อื่น ใส่ Order เป็น 0 ตามไฟล์ดิบ (1 แถวต่อ 1 Invoice)
    no_so = no_so.copy()
    no_so['OrderNo'] = 0
    o, g1 = fill(pd.concat([orders, no_so], ignore_index=True))
    n, g2 = fill(no_so.iloc[0:0])

    # เรียงตามวันที่ -> Warehouse -> ชื่อร้าน (ร้านเดียวกันในวันเดียวกันอยู่ติดกัน)
    o = o.assign(_dk=o['Order Date'].map(lambda v: '' if v is None else str(v)),
                 _wk=o['Warehouse'].map(lambda v: '~' if _cell_blank(v) else str(v)),
                 _sk=o['_key'].map(lambda v: v or '')).sort_values(['_dk', '_wk', '_sk'], kind='stable')

    # Service Time: เวลามาตรฐานตามประเภท หารด้วยจำนวน SO ของ ShipTo เดียวกัน "ในวันเดียวกัน"
    # (ไฟล์ภูเก็ตมีหลายวันในไฟล์เดียว ร้านเดียวกันคนละวัน = คนละรอบส่ง) นับรวมแถวที่ไม่มี SO ด้วย
    std = {'F': float(fs_minutes), 'R': float(retail_minutes)}
    both = pd.concat([o[['Order Date', '_key']], n[['Order Date', '_key']]], ignore_index=True)
    counts = both.dropna(subset=['_key']).groupby(['Order Date', '_key'], dropna=False).size().to_dict()

    def svc(df):
        out = []
        for d, k, t in zip(df['Order Date'], df['_key'], df['Cust Type']):
            if t not in std:
                out.append(None)
            elif not k:
                out.append(std[t])
            else:
                out.append(round(std[t] / counts.get((d, k), 1), 2))
        return out
    o[PKT_SERVICE_COL] = svc(o)
    n[PKT_SERVICE_COL] = svc(n)

    o_out = o[PKT_COLUMNS + ['Term', 'Legal Number', 'Total Amount', '_date']].reset_index(drop=True)
    n_out = n[PKT_COLUMNS].reset_index(drop=True)
    is_no_so = o_out['OrderNo'] == 0
    info = {
        'orders': int((~is_no_so).sum()), 'no_so': int(is_no_so.sum()), 'rows': len(o_out),
        'coords': int(o_out['Latitude'].notna().sum()), 'coords_no_so': int(o_out.loc[is_no_so, 'Latitude'].notna().sum()),
        'type_f': int((o_out['Cust Type'] == 'F').sum()), 'type_r': int((o_out['Cust Type'] == 'R').sum()),
        'type_blank': int(o_out['Cust Type'].isna().sum()), 'type_by_group': g1 + g2,
        'unmatched': sorted(set(
            (r['ShipTo Name'] or r['Customer Name'] or '') for _, r in pd.concat([o_out, n_out]).iterrows()
            if pd.isna(r['Latitude']))),
    }
    return o_out, n_out, info


def pkt_to_excel_bytes(orders_df, no_so_df, sheet1_df=None):
    """Sheet1 ต้นฉบับ (ตัดแถว/คอลัมน์ว่าง) + ชีท Order (เทมเพลตใบงานภูเก็ต + Service Time + Lat/Lon)"""
    wb = Workbook()
    wb.remove(wb.active)
    thin = Side(style='thin')
    border = Border(top=thin, bottom=thin, left=thin, right=thin)
    f_head = Font(name='Cordia New', size=14, bold=True)
    f_body = Font(name='Cordia New', size=14)
    if sheet1_df is not None and len(sheet1_df):
        ws1 = wb.create_sheet(title='Sheet1')
        for r_i, row in enumerate(sheet1_df.itertuples(index=False), start=1):
            for c_i, v in enumerate(row, start=1):
                if _cell_blank(v):
                    continue
                if isinstance(v, (pd.Timestamp, datetime)):
                    v = v.strftime('%Y/%m/%d')
                cell = ws1.cell(row=r_i, column=c_i, value=v)
                cell.font = f_head if r_i == 1 else f_body
        for c_i, col in enumerate(sheet1_df.columns, start=1):
            width = max((len(str(v)) for v in sheet1_df[col] if not _cell_blank(v)), default=6)
            ws1.column_dimensions[get_column_letter(c_i)].width = min(max(width + 2, 8), 50)
    # ---- ชีท Order = ใบงานตามเทมเพลตภูเก็ต: หัวกระดาษ 2 แถว, หัวตารางแถว 3, Route ว่าง,
    #      ต่อท้ายด้วย Service Time + Latitude + Longitude, ท้าย "รายชื่อ พขร." ----
    ws = wb.create_sheet(title=PKT_JOB_SHEET)
    f_top = Font(name='Cordia New', size=16, bold=True)
    f_cell = Font(name='Aptos Narrow', size=11)       # ตารางใช้ฟอนต์เดียวกับเทมเพลตภูเก็ต
    f_hcell = Font(name='Aptos Narrow', size=11, bold=True)
    for r, text in enumerate(PKT_JOB_TOP, start=1):
        c = ws.cell(row=r, column=1, value=text)
        c.font = f_top
        c.alignment = Alignment(horizontal='left', vertical='center')
        ws.row_dimensions[r].height = 24.6
    for c, name in enumerate(PKT_JOB_COLUMNS, start=1):
        cell = ws.cell(row=3, column=c, value=name)
        cell.font = f_hcell
        cell.border = border
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    r = 3
    for rec in orders_df.itertuples(index=False):
        row_vals = dict(zip(orders_df.columns, rec))
        d = row_vals.get('_date')
        if isinstance(d, pd.Timestamp):
            d = d.to_pydatetime()
        amt = row_vals.get('Total Amount')
        vals = [d, row_vals.get('OrderNo'), row_vals.get('Customer Name'), row_vals.get('ShipTo Name'),
                row_vals.get('Term'), row_vals.get('Sales Person'), row_vals.get('Legal Number'),
                None if (amt is None or (isinstance(amt, float) and pd.isna(amt))) else amt,
                None, None, None, None,
                row_vals.get(PKT_SERVICE_COL), row_vals.get('Latitude'), row_vals.get('Longitude')]
        r += 1
        for c, v in enumerate(vals, start=1):
            if v is not None and isinstance(v, float) and pd.isna(v):
                v = None
            cell = ws.cell(row=r, column=c, value=v)
            cell.font = f_cell
            cell.border = border
            cell.alignment = Alignment(vertical='center')
            if c == 1 and isinstance(v, datetime):
                cell.number_format = 'dd/mm/yyyy'
            elif c == 8:
                cell.number_format = '#,##0.00'
    foot = r + 3
    ws.merge_cells(start_row=foot, start_column=3, end_row=foot, end_column=11)
    fc = ws.cell(row=foot, column=3, value=PKT_JOB_FOOTER)
    fc.font = f_hcell
    fc.alignment = Alignment(horizontal='left', vertical='center')
    for c in range(3, 12):
        ws.cell(row=foot, column=c).border = border
    ws.row_dimensions[foot].height = 30
    for c, w in enumerate(PKT_JOB_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.print_title_rows = '1:3'
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _render_pkt():
    if not MASTER_SHEET_CSV_URL:
        st.error("⚠️ ยังไม่ได้ตั้งค่า MASTER_SHEET_CSV_URL — ดูวิธีตั้งค่าใน DEPLOY_INSTRUCTIONS.md")
        return
    st.markdown("##### 🏝️ For PKT — ทำใบงานจากไฟล์ดิบภูเก็ต (1 แถวต่อ 1 SO)")
    pkt_file = st.file_uploader("อัพโหลดไฟล์ดิบภูเก็ต", type=['xls', 'xlsx'], label_visibility="collapsed", key="pkt_uploader")
    if pkt_file is None:
        return
    pkt_bytes = pkt_file.read()
    try:
        pkt_orders, pkt_no_so, pkt_info = read_pkt_orders(pkt_bytes)
    except Exception as e:
        st.error(f"อ่านไฟล์ไม่ได้: {e}")
        return

    st.markdown(f"""
    <div class="rcm-card">
        📄 ชีท <b>{pkt_info['sheet']}</b>: {pkt_info['lines']:,} Invoice → <b>{pkt_info['orders']:,} SO</b> (ไม่ซ้ำ)
        {f" · ไม่มีเลข SO {pkt_info['no_so']:,} Invoice" if pkt_info['no_so'] else ''}
    </div>
    """, unsafe_allow_html=True)

    pkt_dates = sorted({str(d) for d in pd.concat([pkt_orders['Order Date'], pkt_no_so['Order Date']]) if d is not None})
    if len(pkt_dates) > 1:
        pkt_sel = st.multiselect("Inv Date ที่จะเอา (ค่าเริ่มต้น = ทุกวันที่มีในไฟล์)", options=pkt_dates,
                                 default=pkt_dates, key="pkt_dates")
        if not pkt_sel:
            st.warning("⚠️ กรุณาเลือก Inv Date อย่างน้อย 1 วัน")
            return
    else:
        pkt_sel = None

    st.caption("⏱️ เวลาบริการมาตรฐานต่อร้าน (นาที) — หารเท่า ๆ กันตามจำนวน SO ที่ ShipTo Name เดียวกัน")
    s1, s2 = st.columns(2)
    with s1:
        pkt_fs = st.number_input("Food Service", min_value=1.0, max_value=240.0, value=20.0, step=1.0, key="pkt_service_fs")
    with s2:
        pkt_rt = st.number_input("Retail", min_value=1.0, max_value=240.0, value=45.0, step=1.0, key="pkt_service_rt")

    if st.button("🚀 ทำใบงาน For PKT", type="primary", use_container_width=True, key="pkt_btn"):
        with st.spinner("🛰️ กำลังรวม SO และเติมพิกัด..."):
            try:
                master_df = load_master_data(MASTER_SHEET_CSV_URL)
                channels = load_master_channels(MASTER_SHEET_CSV_URL)
                o, n, res = process_pkt_orders(pkt_orders, pkt_no_so, master_df, channels,
                                               dates=pkt_sel, fs_minutes=pkt_fs, retail_minutes=pkt_rt)
                st.session_state['pkt_xlsx'] = pkt_to_excel_bytes(o, n, sheet1_df=bkk_sheet1_clean(pkt_bytes))
                st.session_state['pkt_result'] = res
                st.session_state['pkt_filename'] = pkt_file.name
            except Exception as e:
                st.error(f"เกิดข้อผิดพลาด: {e}")
                st.exception(e)

    if 'pkt_result' not in st.session_state or st.session_state.get('pkt_filename') != pkt_file.name:
        return
    res = st.session_state['pkt_result']
    st.markdown(f"""
    <div class="rcm-card rcm-card-success">
        🎯 <b>เสร็จแล้ว!</b> {res['orders']:,} SO + ไม่มี SO {res['no_so']:,} · มีพิกัด {res['coords']:,} / {res['rows']:,}
    </div>
    """, unsafe_allow_html=True)
    c1, c2, c3 = st.columns(3)
    with c1:
        stat_box("SO", f"{res['orders']:,}", f"ไม่ซ้ำ · + Invoice ที่ไม่มี SO {res['no_so']:,}")
    with c2:
        stat_box("มีพิกัด", f"{res['coords']:,} / {res['rows']:,}", "แถวทั้งหมดในชีท Order")
    with c3:
        stat_box("CUST TYPE", f"F {res['type_f']:,} · R {res['type_r']:,}",
                 f"ไม่ทราบประเภท {res['type_blank']:,} (Service Time ว่าง)")
    notes = []
    if res['no_so']:
        notes.append(f"{res['no_so']:,} Invoice ไม่มีเลข SO (Order = 0) — เก็บไว้ในชีท Order ด้วย 1 แถวต่อ 1 Invoice")
    if res['type_blank']:
        notes.append(f"{res['type_blank']:,} SO ไม่ทราบว่าเป็น F หรือ R (Channel ใน Master ไม่ใช่ Food Service/Retail และ Group ในไฟล์ก็ไม่บอก) — ช่อง Service Time ว่างไว้ให้กรอกเอง")
    if notes:
        st.markdown('<div class="rcm-card rcm-card-warn">⚠️ <b>ข้อสังเกต</b><br>' + '<br>'.join('• ' + x for x in notes) + '</div>',
                    unsafe_allow_html=True)
    if res['unmatched']:
        with st.expander(f"ร้านที่ยังไม่มีพิกัด ({len(res['unmatched']):,} ชื่อ) — เพิ่มใน Master แล้วกดทำใหม่"):
            st.dataframe(pd.DataFrame({'ShipTo Name / Customer Name': res['unmatched']}),
                         use_container_width=True, hide_index=True)
    st.write("")
    st.download_button(
        "⬇️ ดาวน์โหลดใบงาน For PKT (Excel: Sheet1 + Order)",
        data=st.session_state['pkt_xlsx'],
        file_name=f"{pkt_file.name.rsplit('.', 1)[0]}_For_PKT.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        use_container_width=True, key="pkt_download_btn",
    )


# ============================================================
# แท็บของแต่ละพื้นที่
#   BKK   : For BKK + ใส่พิกัด Lat/Lon + Weight Calculator (+ QR Code)
#   PKT   : For PKT + ใส่พิกัด Lat/Lon (+ QR Code)
#   SAMUI : ใส่พิกัด Lat/Lon (+ QR Code)
# ============================================================
REGION_TABS = {
    'BKK': [("🏙️  For BKK", _render_bkk), ("📍  ใส่พิกัด Lat/Lon", _render_coords),
            ("⚖️  Weight Calculator", _render_weight), ("🔗  QR Code", _render_qr)],
    'PKT': [("🏝️  For PKT", _render_pkt), ("📍  ใส่พิกัด Lat/Lon", _render_coords), ("🔗  QR Code", _render_qr)],
    'SAMUI': [("📍  ใส่พิกัด Lat/Lon", _render_coords), ("🔗  QR Code", _render_qr)],
}

_region_tabs = REGION_TABS.get(region, REGION_TABS['BKK'])
for _tab, (_label, _render) in zip(st.tabs([label for label, _ in _region_tabs]), _region_tabs):
    with _tab:
        _render()
